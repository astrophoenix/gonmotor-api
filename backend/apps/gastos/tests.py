from datetime import date
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.test import TestCase
from rest_framework.test import APIClient

from apps.empresas.models import Empresa, Taller

from .models import Gasto
from apps.core.testing import fijar_empresa

User = get_user_model()
ENDPOINT = '/api/gastos/'


class GastosFiltroTallerTests(TestCase):
    """El listado se recorta al taller activo (`X-Taller-ID`), como los demás módulos."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser(username='gastos_admin', password='x')
        cls.empresa = Empresa.objects.create(
            nombre_comercial='Taller Gastos',
            ruc=str(uuid4().int % 10_000_000_000_000).zfill(13),
            email_contacto='gastos@test.com',
        )
        cls.central = Taller.objects.create(
            empresa=cls.empresa, nombre='Central', codigo_sucursal='001', direccion='Av. 1',
        )
        cls.norte = Taller.objects.create(
            empresa=cls.empresa, nombre='Norte', codigo_sucursal='002', direccion='Av. 2',
        )
        cls.gasto_central = Gasto.objects.create(
            empresa=cls.empresa, taller=cls.central, fecha=date(2026, 1, 10), monto=10,
        )
        cls.gasto_norte = Gasto.objects.create(
            empresa=cls.empresa, taller=cls.norte, fecha=date(2026, 1, 11), monto=20,
        )

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        fijar_empresa(self.client, str(self.empresa.uuid))
    def _ids(self, respuesta):
        self.assertEqual(respuesta.status_code, 200)
        return {fila['id'] for fila in respuesta.data['results']}

    def test_con_taller_activo_solo_lista_sus_gastos(self):
        respuesta = self.client.get(ENDPOINT, HTTP_X_TALLER_ID=str(self.central.uuid))
        self.assertEqual(self._ids(respuesta), {str(self.gasto_central.uuid)})

    def test_todos_los_talleres_lista_todo(self):
        respuesta = self.client.get(ENDPOINT)
        self.assertEqual(self._ids(respuesta), {str(self.gasto_central.uuid), str(self.gasto_norte.uuid)})

    def test_filtro_manual_por_taller_en_todos(self):
        respuesta = self.client.get(ENDPOINT, {'taller_id': str(self.norte.uuid)})
        self.assertEqual(self._ids(respuesta), {str(self.gasto_norte.uuid)})


class GastosCrearTests(TestCase):
    """Alta de gastos desde el modal: la empresa sale de la sesión y el taller,
    del selector visible o del taller activo (`X-Taller-ID`)."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser(username='gastos_crear', password='x')
        cls.empresa = Empresa.objects.create(
            nombre_comercial='Taller Crear Gastos',
            ruc=str(uuid4().int % 10_000_000_000_000).zfill(13),
            email_contacto='crear@test.com',
        )
        cls.otra_empresa = Empresa.objects.create(
            nombre_comercial='Ajena',
            ruc=str(uuid4().int % 10_000_000_000_000).zfill(13),
            email_contacto='ajena@test.com',
        )
        cls.central = Taller.objects.create(
            empresa=cls.empresa, nombre='Central', codigo_sucursal='001', direccion='Av. 1',
        )
        cls.norte = Taller.objects.create(
            empresa=cls.empresa, nombre='Norte', codigo_sucursal='002', direccion='Av. 2',
        )
        cls.taller_ajeno = Taller.objects.create(
            empresa=cls.otra_empresa, nombre='Ajeno', codigo_sucursal='001', direccion='X',
        )

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        fijar_empresa(self.client, str(self.empresa.uuid))
    def _payload(self, **extra):
        return {
            'categoria': 'INSUMOS', 'fecha': '2026-10-14', 'monto': 20,
            'descripcion': 'gasto pagado por transferencia', **extra,
        }

    def test_crear_sin_empresa_ni_taller_usa_la_sesion(self):
        respuesta = self.client.post(
            ENDPOINT, self._payload(), format='json', HTTP_X_TALLER_ID=str(self.norte.uuid),
        )
        self.assertEqual(respuesta.status_code, 201, respuesta.content)
        gasto = Gasto.objects.get(uuid=respuesta.data['id'])
        self.assertEqual(gasto.empresa_id, self.empresa.pk)
        self.assertEqual(gasto.taller_id, self.norte.pk)

    def test_crear_con_taller_explicito(self):
        respuesta = self.client.post(ENDPOINT, self._payload(taller=str(self.central.uuid)), format='json')
        self.assertEqual(respuesta.status_code, 201, respuesta.content)
        self.assertEqual(Gasto.objects.get(uuid=respuesta.data['id']).taller_id, self.central.pk)

    def test_no_se_puede_forzar_otra_empresa(self):
        respuesta = self.client.post(ENDPOINT, self._payload(empresa=self.otra_empresa.pk), format='json')
        self.assertEqual(respuesta.status_code, 201, respuesta.content)
        self.assertEqual(Gasto.objects.get(uuid=respuesta.data['id']).empresa_id, self.empresa.pk)

    def test_taller_de_otra_empresa_se_rechaza(self):
        respuesta = self.client.post(ENDPOINT, self._payload(taller=str(self.taller_ajeno.uuid)), format='json')
        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('taller', respuesta.data)


class GastosExportarExcelTests(TestCase):
    """El Excel usa los mismos filtros y taller que el listado."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser(username='gastos_excel', password='x')
        cls.empresa = Empresa.objects.create(
            nombre_comercial='Taller Excel Gastos',
            ruc=str(uuid4().int % 10_000_000_000_000).zfill(13),
            email_contacto='excel@test.com',
        )
        cls.central = Taller.objects.create(
            empresa=cls.empresa, nombre='Central', codigo_sucursal='001', direccion='Av. 1',
        )
        cls.norte = Taller.objects.create(
            empresa=cls.empresa, nombre='Norte', codigo_sucursal='002', direccion='Av. 2',
        )
        Gasto.objects.create(
            empresa=cls.empresa, taller=cls.central, fecha=date(2026, 1, 10), monto=10,
            categoria='ARRIENDO', descripcion='=1+1',
        )
        Gasto.objects.create(
            empresa=cls.empresa, taller=cls.norte, fecha=date(2026, 2, 11), monto=20, categoria='INSUMOS',
        )

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        fijar_empresa(self.client, str(self.empresa.uuid))

    def _filas(self, **params):
        from apps.pagos.tests_exportes import filas_excel

        respuesta = self.client.get(f'{ENDPOINT}export-excel/', params, **params.pop('headers', {}))
        self.assertEqual(respuesta.status_code, 200)
        return filas_excel(respuesta, 'Fecha')

    def test_exporta_con_filtros(self):
        self.assertEqual(len(self._filas()), 2)
        self.assertEqual([fila[2] for fila in self._filas(categoria='INSUMOS')], ['Insumos/Compras menores'])
        self.assertEqual([fila[0] for fila in self._filas(fecha_desde='2026-02-01')], ['11/02/2026'])

    def test_taller_activo_y_formulas_neutralizadas(self):
        respuesta = self.client.get(f'{ENDPOINT}export-excel/', HTTP_X_TALLER_ID=str(self.central.uuid))
        from io import BytesIO

        from openpyxl import load_workbook

        hoja = load_workbook(BytesIO(respuesta.content)).active
        celdas = [celda for fila in hoja.iter_rows() for celda in fila if celda.value == '=1+1']
        self.assertEqual(len(celdas), 1)
        self.assertEqual(celdas[0].data_type, 's')
        textos = [celda.value for fila in hoja.iter_rows() for celda in fila]
        self.assertNotIn('Norte', textos)
