from io import BytesIO
from unittest.mock import patch
from openpyxl import load_workbook

from datetime import date, datetime, time
from decimal import Decimal
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from apps.authentication.models import UsuarioEmpresa
from apps.clientes.models import Cliente
from apps.core.testing import fijar_empresa
from apps.empresas.models import Empresa, Taller
from apps.gastos.models import Gasto
from apps.ordenes.models import DetalleRepuestoOrdenTrabajo, DetalleServicioOrdenTrabajo, OrdenTrabajo
from apps.pagos.models import PagoOrden
from apps.vehiculos.models import Vehiculo, vincular_empresa

User = get_user_model()
URL = '/api/reportes/resumen/'
# Periodo de prueba: junio 2026; el anterior de igual duración es mayo 2–31.
JUNIO = {'fecha_desde': '2026-06-01', 'fecha_hasta': '2026-06-30'}


def _ruc():
    return str(uuid4().int % 10_000_000_000_000).zfill(13)


def _momento(dia, hora=10):
    return timezone.make_aware(datetime.combine(dia, time(hora, 0)))


class ResumenNegocioTests(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.empresa = Empresa.objects.create(nombre_comercial='Taller Resumen', ruc=_ruc())
        cls.otra = Empresa.objects.create(nombre_comercial='Ajena', ruc=_ruc())
        cls.a1 = Taller.objects.create(empresa=cls.empresa, nombre='A1', codigo_sucursal='001', direccion='x')
        cls.a2 = Taller.objects.create(empresa=cls.empresa, nombre='A2', codigo_sucursal='002', direccion='x')
        cls.b1 = Taller.objects.create(empresa=cls.otra, nombre='B1', codigo_sucursal='001', direccion='x')
        cls.cliente = Cliente.objects.create(empresa=cls.empresa, nombre='Cliente', identificacion='1700000021')
        cls.cliente_b = Cliente.objects.create(empresa=cls.otra, nombre='Ajeno', identificacion='1700000022')
        cls.vehiculo = Vehiculo.objects.create(placa='RES0001', marca='Kia', modelo='Rio')
        vincular_empresa(cls.vehiculo, cls.empresa.pk)
        vincular_empresa(cls.vehiculo, cls.otra.pk)
        cls.admin = cls._usuario('res_admin', cls.empresa, 'ADMIN_EMPRESA')
        cls.jefe_a1 = cls._usuario('res_jefe_a1', cls.empresa, 'ADMIN_TALLER', [cls.a1])
        cls.mecanico = cls._usuario('res_mecanico', cls.empresa, 'MECANICO', [cls.a1])

        # Junio, A1: OT entregada de $115 (100 + IVA), costo 40; cobró 100 en efectivo y devolvió 10.
        cls.ot_junio = cls._orden('A1-1', cls.a1, entregada=date(2026, 6, 10), creada=date(2026, 6, 5))
        cls._pago(cls.ot_junio, '100', date(2026, 6, 10))
        cls._pago(cls.ot_junio, '10', date(2026, 6, 12), tipo='DEVOLUCION')
        cls._pago(cls.ot_junio, '50', date(2026, 6, 11), anulado=True)
        # Junio, A2: OT entregada de $115 pagada por transferencia.
        cls.ot_a2 = cls._orden('A2-1', cls.a2, entregada=date(2026, 6, 20), creada=date(2026, 6, 20))
        cls._pago(cls.ot_a2, '115', date(2026, 6, 20), forma='TRANSFERENCIA')
        # Mayo, A1: OT entregada de $115 sin pagos (periodo anterior).
        cls.ot_mayo = cls._orden('A1-0', cls.a1, entregada=date(2026, 5, 15), creada=date(2026, 5, 10))
        # Julio (fuera del periodo) y una abierta creada en junio.
        cls._orden('A1-2', cls.a1, entregada=date(2026, 7, 2), creada=date(2026, 6, 28))
        cls._orden('A1-3', cls.a1, creada=date(2026, 6, 15))
        # Otra empresa: no debe sumar.
        cls._orden('B-1', cls.b1, entregada=date(2026, 6, 10), creada=date(2026, 6, 1), empresa=cls.otra, cliente=cls.cliente_b)

        Gasto.objects.create(empresa=cls.empresa, taller=cls.a1, fecha=date(2026, 6, 3), categoria='ARRIENDO', monto=30)
        Gasto.objects.create(empresa=cls.empresa, taller=cls.a2, fecha=date(2026, 6, 4), categoria='INSUMOS', monto=5)
        Gasto.objects.create(empresa=cls.empresa, taller=cls.a1, fecha=date(2026, 5, 20), categoria='ARRIENDO', monto=30)

    @staticmethod
    def _usuario(username, empresa, rol, talleres=None):
        user = User.objects.create_user(username=username, password='x', email=f'{username}@test.ec')
        relacion = UsuarioEmpresa.objects.create(user=user, empresa=empresa, rol=rol)
        if talleres:
            relacion.talleres.set(talleres)
        return user

    @classmethod
    def _orden(cls, numero, taller, *, creada, entregada=None, empresa=None, cliente=None):
        orden = OrdenTrabajo.objects.create(
            empresa=empresa or cls.empresa, sucursal=taller, cliente=cliente or cls.cliente,
            vehiculo=cls.vehiculo, numero_orden=numero,
        )
        DetalleRepuestoOrdenTrabajo.objects.create(
            orden_trabajo=orden, descripcion='Filtro', precio_unitario=Decimal('40'), costo_unitario=Decimal('25'),
        )
        DetalleServicioOrdenTrabajo.objects.create(
            orden_trabajo=orden, descripcion='Cambio de aceite', precio_unitario=Decimal('60'),
            costo_unitario=Decimal('15'),
        )
        orden.calcular_totales()
        campos = {'created_at': _momento(creada)}
        if entregada:
            campos.update(estado='ENTREGADO', fecha_entrega=_momento(entregada))
        OrdenTrabajo.objects.filter(pk=orden.pk).update(**campos)
        return orden

    @staticmethod
    def _pago(orden, monto, dia, *, tipo='PAGO', forma='EFECTIVO', anulado=False):
        return PagoOrden.objects.create(
            orden_trabajo=orden, tipo=tipo, forma_pago=forma, monto=Decimal(monto),
            fecha=_momento(dia), anulado=anulado,
        )

    def _api(self, user=None):
        client = APIClient()
        client.force_authenticate(user=user or self.admin)
        fijar_empresa(client, str(self.empresa.uuid))
        return client

    def _get(self, params=None, user=None):
        respuesta = self._api(user).get(URL, params or JUNIO)
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        return respuesta.data

    def test_facturado_cobrado_y_utilidad_de_la_empresa(self):
        datos = self._get()
        actual = datos['actual']
        self.assertEqual(actual['ordenes']['entregadas'], 2)
        self.assertEqual(actual['facturado']['total'], Decimal('230.00'))
        self.assertEqual(actual['facturado']['base_gravada'], Decimal('200.00'))
        self.assertEqual(actual['facturado']['iva'], Decimal('30.00'))
        self.assertEqual(actual['ingresos'], Decimal('200.00'))
        # 100 + 115 − 10; el pago anulado no cuenta.
        self.assertEqual(actual['cobrado']['total'], Decimal('205.00'))
        formas = {f['forma_pago']: f['neto'] for f in actual['cobrado']['por_forma']}
        self.assertEqual(formas, {'EFECTIVO': Decimal('90.00'), 'TRANSFERENCIA': Decimal('115.00')})
        self.assertEqual(actual['costo_trabajos']['total'], Decimal('80.00'))
        self.assertEqual(actual['gastos']['total'], Decimal('35.00'))
        self.assertEqual(actual['utilidad'], Decimal('85.00'))
        self.assertEqual(actual['ticket_promedio'], Decimal('115.00'))
        # Al 30/06 quedan por cobrar los 25 de A1-1 (115 − 90) y los 115 de mayo.
        self.assertEqual(actual['por_cobrar']['total'], Decimal('140.00'))
        self.assertEqual(actual['por_cobrar']['ordenes'], 2)

    def test_conteo_de_ordenes(self):
        ordenes = self._get()['actual']['ordenes']
        # Creadas en junio: A1-1, A2-1, A1-2 y A1-3.
        self.assertEqual(ordenes['creadas'], 4)
        # Abiertas al 30/06: A1-3 (en proceso) y A1-2 (se entregó en julio).
        self.assertEqual(ordenes['abiertas'], 2)

    def test_periodo_anterior_y_variacion(self):
        datos = self._get()
        self.assertEqual(datos['periodo_anterior'], {'desde': '2026-05-02', 'hasta': '2026-05-31'})
        facturado = datos['kpis']['facturado']
        self.assertEqual(facturado['anterior'], Decimal('115.00'))
        self.assertEqual(facturado['variacion'], Decimal('115.00'))
        self.assertEqual(facturado['variacion_pct'], 100.0)
        self.assertIsNone(datos['kpis']['cobrado']['variacion_pct'])

    def test_serie_diaria_y_mensual(self):
        serie = self._get()['serie']
        self.assertEqual(serie['agrupacion'], 'dia')
        self.assertEqual(len(serie['puntos']), 30)
        dia_10 = next(p for p in serie['puntos'] if p['fecha'] == '2026-06-10')
        self.assertEqual(dia_10['facturado'], Decimal('115.00'))
        self.assertEqual(dia_10['cobrado'], Decimal('100.00'))
        mensual = self._get({'fecha_desde': '2026-01-01', 'fecha_hasta': '2026-06-30'})['serie']
        self.assertEqual(mensual['agrupacion'], 'mes')
        self.assertEqual([p['fecha'] for p in mensual['puntos']][-1], '2026-06-01')
        self.assertEqual(mensual['puntos'][-1]['facturado'], Decimal('230.00'))

    def test_top_servicios(self):
        top = self._get()['top_servicios']
        self.assertEqual(top[0]['descripcion'], 'Cambio de aceite')
        self.assertEqual(top[0]['veces'], 2)

    def test_recorte_por_taller(self):
        jefe = self._get(user=self.jefe_a1)['actual']
        self.assertEqual(jefe['facturado']['total'], Decimal('115.00'))
        self.assertEqual(jefe['gastos']['total'], Decimal('30.00'))
        admin_a2 = self._get({**JUNIO, 'taller_id': str(self.a2.uuid)})['actual']
        self.assertEqual(admin_a2['cobrado']['total'], Decimal('115.00'))

    def test_mecanico_sin_reportes_no_entra(self):
        self.assertEqual(self._api(self.mecanico).get(URL, JUNIO).status_code, 403)

    def test_validacion_de_fechas(self):
        api = self._api()
        self.assertEqual(api.get(URL, {'fecha_desde': '2026-06-30', 'fecha_hasta': '2026-06-01'}).status_code, 400)
        self.assertEqual(api.get(URL, {'fecha_desde': '2025-01-01', 'fecha_hasta': '2026-06-30'}).status_code, 400)
        self.assertEqual(api.get(URL, {'fecha_desde': 'junio'}).status_code, 400)
        # Sin fechas: el mes en curso hasta hoy.
        periodo = api.get(URL).data['periodo']
        self.assertEqual(periodo['hasta'], timezone.localdate().isoformat())
        self.assertTrue(periodo['desde'].endswith('-01'))

    def test_pdf(self):
        respuesta = self._api().get(f'{URL}exportar-pdf/', JUNIO)
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')

    def test_excel_mismos_indicadores_numericos_y_periodos(self):
        r = self._api().get(f'{URL}export-excel/', JUNIO)
        self.assertEqual(r.status_code, 200, r.content)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('2026-06-01_2026-06-30.xlsx', r['Content-Disposition'])
        ws = load_workbook(BytesIO(r.content)).active
        filas = {fila[0]: fila[1:] for fila in ws.iter_rows(values_only=True) if fila[0]}
        self.assertEqual(filas['Facturado (con IVA)'][:4], (230, 115, 115, 100))
        self.assertEqual(filas['Utilidad'][0], 85)
        self.assertEqual(filas['Cobrado en efectivo'][0], 90)
        self.assertEqual(filas['Órdenes entregadas'][0], 2)
        self.assertTrue(any('2026-05-02' in str(v) for row in ws.values for v in row))
        for row in ws.iter_rows():
            if row[0].value == 'Facturado (con IVA)':
                self.assertEqual(row[1].data_type, 'n')

    def test_excel_respeta_taller_tenant_permiso_y_fechas(self):
        r = self._api(self.jefe_a1).get(f'{URL}export-excel/', JUNIO)
        ws = load_workbook(BytesIO(r.content)).active
        filas = {fila[0]: fila[1:] for fila in ws.iter_rows(values_only=True) if fila[0]}
        self.assertEqual(filas['Facturado (con IVA)'][0], 115)
        self.assertEqual(self._api(self.mecanico).get(f'{URL}export-excel/', JUNIO).status_code, 403)
        self.assertEqual(self._api().get(f'{URL}export-excel/', {'fecha_desde':'2026-99-01'}).status_code, 400)
        api = self._api(); fijar_empresa(api, self.otra)
        self.assertEqual(api.get(f'{URL}export-excel/', JUNIO).status_code, 403)

    def test_excel_escapa_nombre_de_empresa_como_texto(self):
        self.empresa.nombre_comercial = '=HYPERLINK("https://example.test","x")'
        self.empresa.save()
        r = self._api().get(f'{URL}export-excel/', JUNIO)
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual(ws['A1'].data_type, 's')
        self.assertTrue(ws['A1'].value.startswith('='))

    def test_fecha_anulacion_fija_no_cambia_con_edicion_y_no_es_escribible(self):
        orden = self._orden('CANCEL-TEST', self.a1, creada=date(2026,6,1))
        anula = _momento(date(2026,6,10))
        with patch('django.utils.timezone.now', return_value=anula):
            r = self._api().patch(f'/api/ordenes/ordenes-trabajo/{orden.uuid}/', {'estado':'CANCELADO','fecha_anulacion':'2020-01-01T00:00:00Z'}, format='json')
        self.assertEqual(r.status_code,200,r.data)
        orden.refresh_from_db(); self.assertEqual(orden.fecha_anulacion,anula)
        with patch('django.utils.timezone.now', return_value=_momento(date(2026,7,5))):
            r = self._api().patch(f'/api/ordenes/ordenes-trabajo/{orden.uuid}/', {'observaciones':'Corrección'}, format='json')
        self.assertEqual(r.status_code,200,r.data)
        orden.refresh_from_db(); self.assertEqual(orden.fecha_anulacion,anula)
        self.assertEqual(self._get()['actual']['ordenes']['anuladas'],1)
        self.assertEqual(self._get({'fecha_desde':'2026-07-01','fecha_hasta':'2026-07-31'})['actual']['ordenes']['anuladas'],0)

    def test_anulacion_modelo_y_update_fields_conservan_fecha(self):
        orden = self._orden('CANCEL-MODEL', self.a1, creada=date(2026,6,1))
        anula = _momento(date(2026,6,10))
        with patch('django.utils.timezone.now', return_value=anula):
            orden.estado = 'CANCELADO'; orden.save(update_fields=['estado'])
        orden.refresh_from_db(); self.assertEqual(orden.fecha_anulacion,anula)
        orden.observaciones = 'Nota'; orden.save(update_fields=['observaciones'])
        orden.refresh_from_db(); self.assertEqual(orden.fecha_anulacion,anula)
