from django.contrib.auth.models import User
from rest_framework import status
from rest_framework.test import APIClient, APITestCase


class RegistrationTests(APITestCase):
	endpoint = '/api/auth/register/'

	def valid_payload(self, **overrides):
		payload = {
			'email': 'nuevo@ejemplo.com',
			'password': 'UnaClaveSegura123!',
			'password_confirmation': 'UnaClaveSegura123!',
			'accepted_terms': True,
		}
		payload.update(overrides)
		return payload

	def test_registers_user_with_email_as_username(self):
		response = self.client.post(self.endpoint, self.valid_payload(), format='json')

		self.assertEqual(response.status_code, status.HTTP_201_CREATED)
		self.assertTrue(User.objects.filter(username='nuevo@ejemplo.com').exists())

	def test_rejects_mismatched_passwords(self):
		response = self.client.post(
			self.endpoint,
			self.valid_payload(password_confirmation='OtraClave123!'),
			format='json'
		)

		self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
		self.assertIn('password_confirmation', response.data)

	def test_rejects_duplicate_email(self):
		User.objects.create_user(
			username='existente@ejemplo.com',
			email='existente@ejemplo.com',
			password='UnaClaveSegura123!'
		)

		response = self.client.post(
			self.endpoint,
			self.valid_payload(email='EXISTENTE@EJEMPLO.COM'),
			format='json'
		)

		self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
		self.assertIn('email', response.data)


class EmpleadoExportTests(APITestCase):
	def setUp(self):
		from apps.empresas.models import Empresa

		self.user = User.objects.create_superuser(username='export_tester', password='x')
		self.empresa = Empresa.objects.create(
			nombre_comercial='Taller GON',
			razon_social='GON SA',
			ruc='1750000000004',
		)
		self.usuario = User.objects.create_user(
			username='empleado1',
			email='empleado1@ejemplo.com',
			password='UnaClaveSegura123!',
			first_name='Ana',
			last_name='Gómez',
		)
		self.client = APIClient()
		self.client.force_authenticate(user=self.user)
		self.client.defaults['HTTP_X_EMPRESA_ID'] = self.empresa.pk

	def _crear_empleado(self):
		from apps.authentication.models import UsuarioEmpresa

		ue = UsuarioEmpresa.objects.create(
			user=self.usuario,
			empresa=self.empresa,
			rol='ASESOR',
		)
		return ue

	def test_exportar_pdf_empleados(self):
		self._crear_empleado()
		resp = self.client.get('/api/auth/empleados/exportar-pdf/')
		self.assertEqual(resp.status_code, 200)
		self.assertEqual(resp['Content-Type'], 'application/pdf')
		contenido = b''.join(resp.streaming_content)
		self.assertTrue(contenido.startswith(b'%PDF'))

	def test_exportar_excel_empleados(self):
		self._crear_empleado()
		resp = self.client.get('/api/auth/empleados/export-excel/')
		self.assertEqual(resp.status_code, 200)
		self.assertTrue(resp.content.startswith(b'PK'))

	def test_excel_contiene_datos_empleado(self):
		from io import BytesIO
		from openpyxl import load_workbook

		self._crear_empleado()
		resp = self.client.get('/api/auth/empleados/export-excel/')
		wb = load_workbook(BytesIO(resp.content))
		valores = [
			celda
			for fila in wb.active.iter_rows(values_only=True)
			for celda in fila
		]
		self.assertIn('empleado1@ejemplo.com', valores)
		self.assertIn('Asesor de Servicio', valores)
		self.assertIn('Activo', valores)

	def test_exportar_sin_empresa_rechazado(self):
		self.client.force_authenticate(user=self.usuario)
		self.client.defaults.pop('HTTP_X_EMPRESA_ID')
		resp = self.client.get('/api/auth/empleados/exportar-pdf/')
		self.assertEqual(resp.status_code, 403)


class AvatarUploadTests(APITestCase):
	def setUp(self):
		from django.core.files.uploadedfile import SimpleUploadedFile
		from io import BytesIO
		from PIL import Image

		self.SimpleUploadedFile = SimpleUploadedFile
		self.user = User.objects.create_user(
			username='avatar_user',
			email='avatar@ejemplo.com',
			password='UnaClaveSegura123!',
			first_name='Luis',
			last_name='Pérez',
		)
		self.client = APIClient()
		self.client.force_authenticate(user=self.user)

		def _imagen_png():
			buf = BytesIO()
			Image.new('RGB', (80, 80), color='red').save(buf, format='PNG')
			return buf.getvalue()

		self.png_bytes = _imagen_png()

	def test_subir_avatar_validando_formato(self):
		foto = self.SimpleUploadedFile(
			'foto.png', self.png_bytes, content_type='image/png'
		)
		resp = self.client.patch(
			'/api/auth/me/',
			{'avatar': foto},
			format='multipart',
		)
		self.assertEqual(resp.status_code, 200, resp.data)
		self.assertIn('avatar', resp.data['user'])
		self.assertTrue(resp.data['user']['avatar'])

	def test_rechaza_formato_extraño(self):
		foto = self.SimpleUploadedFile(
			'foto.txt', b'no soy una imagen', content_type='text/plain'
		)
		resp = self.client.patch(
			'/api/auth/me/',
			{'avatar': foto},
			format='multipart',
		)
		self.assertEqual(resp.status_code, 400)
		self.assertIn('avatar', resp.data)

	def test_rechaza_imagen_muy_grande(self):
		from django.core.files.uploadedfile import SimpleUploadedFile
		from io import BytesIO
		from PIL import Image
		import os

		buf = BytesIO()
		ruido = Image.frombytes('RGB', (2200, 2200), os.urandom(2200 * 2200 * 3))
		ruido.save(buf, format='PNG')
		foto = SimpleUploadedFile('grande.png', buf.getvalue(), content_type='image/png')
		self.assertGreater(foto.size, 1024 * 1024)
		resp = self.client.patch(
			'/api/auth/me/',
			{'avatar': foto},
			format='multipart',
		)
		self.assertEqual(resp.status_code, 400)
		self.assertIn('avatar', resp.data)

	def test_eliminar_avatar(self):
		foto = self.SimpleUploadedFile(
			'foto.png', self.png_bytes, content_type='image/png'
		)
		self.client.patch('/api/auth/me/', {'avatar': foto}, format='multipart')

		resp = self.client.patch('/api/auth/me/', {'remove_avatar': 'true'}, format='multipart')
		self.assertEqual(resp.status_code, 200, resp.data)
		self.assertIsNone(resp.data['user']['avatar'])
