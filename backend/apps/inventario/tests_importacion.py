"""Importación del catálogo de inventario: repuestos y servicios."""

from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient, APITestCase

from apps.authentication.models import UsuarioEmpresa
from apps.core.testing import fijar_empresa
from apps.empresas.models import Empresa, Taller
from apps.inventario.excel_import import ImportadorRepuestos, ImportadorServicios
from apps.inventario.models import Repuesto, Servicio

User = get_user_model()


def _xlsx(importador, filas):
    wb = Workbook()
    ws = wb.active
    ws.append([c.encabezado for c in importador.columnas])
    for fila in filas:
        ws.append([fila.get(c.clave, '') for c in importador.columnas])
    buffer = BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile('inventario.xlsx', buffer.getvalue())


class _Base(APITestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre_comercial='Taller Inventario', ruc=str(uuid4().int % 10**13).zfill(13))
        self.taller = Taller.objects.create(empresa=self.empresa, nombre='Central', codigo_sucursal='001', direccion='x')
        self.client = self._como('ADMIN_EMPRESA')

    def _como(self, rol):
        user = User.objects.create_user(username=f'inv-{rol}-{uuid4().hex[:6]}', password='x')
        relacion = UsuarioEmpresa.objects.create(user=user, empresa=self.empresa, rol=rol)
        relacion.talleres.set([self.taller])
        cliente = APIClient()
        cliente.force_authenticate(user=user)
        fijar_empresa(cliente, str(self.empresa.uuid))
        return cliente

    def _importar(self, url, importador, filas, validar=False, cliente=None):
        return (cliente or self.client).post(
            url + ('?validar=1' if validar else ''), {'archivo': _xlsx(importador, filas)}, format='multipart',
        )


class ImportacionRepuestosTests(_Base):
    URL = '/api/repuestos/importar/'

    def _importar_repuestos(self, filas, **kwargs):
        return self._importar(self.URL, ImportadorRepuestos, filas, **kwargs)

    def test_crea_con_valores_por_defecto_y_codigo_en_mayusculas(self):
        respuesta = self._importar_repuestos([{'codigo': 'fil-001', 'nombre': 'Filtro', 'stock_actual': '12,5'}])
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        repuesto = Repuesto.objects.get(empresa=self.empresa, codigo='FIL-001')
        self.assertEqual(repuesto.stock_actual, Decimal('12.50'))
        self.assertEqual(repuesto.categoria, Repuesto.Categoria.OTROS)
        self.assertEqual(repuesto.iva_porcentaje_defecto, Decimal('0.1500'))

    def test_existente_solo_cambian_las_celdas_con_valor(self):
        Repuesto.objects.create(empresa=self.empresa, codigo='FIL-002', nombre='Filtro', marca='Bosch', precio_venta=Decimal('8'))
        self._importar_repuestos([{'codigo': 'FIL-002', 'nombre': 'Filtro de aceite', 'precio_venta': '9.50'}])
        repuesto = Repuesto.objects.get(empresa=self.empresa, codigo='FIL-002')
        self.assertEqual(repuesto.nombre, 'Filtro de aceite')
        self.assertEqual(repuesto.precio_venta, Decimal('9.50'))
        self.assertEqual(repuesto.marca, 'Bosch')

    def test_el_stock_de_un_existente_no_se_pisa_y_se_avisa(self):
        Repuesto.objects.create(empresa=self.empresa, codigo='FIL-003', nombre='Filtro', stock_actual=Decimal('7'))
        previa = self._importar_repuestos([{'codigo': 'FIL-003', 'nombre': 'Filtro', 'stock_actual': '50'}], validar=True).json()
        fila = previa['filas'][0]
        self.assertEqual(fila['estado'], 'actualiza')
        self.assertIn('stock no se cambió', fila['motivo'])
        self.assertEqual(previa['errores'], [])

        self._importar_repuestos([{'codigo': 'FIL-003', 'nombre': 'Filtro', 'stock_actual': '50'}])
        self.assertEqual(Repuesto.objects.get(empresa=self.empresa, codigo='FIL-003').stock_actual, Decimal('7.00'))

    def test_mismo_stock_no_genera_aviso(self):
        Repuesto.objects.create(empresa=self.empresa, codigo='FIL-004', nombre='Filtro', stock_actual=Decimal('7'))
        previa = self._importar_repuestos([{'codigo': 'FIL-004', 'nombre': 'Filtro', 'stock_actual': '7'}], validar=True).json()
        self.assertEqual(previa['filas'][0]['motivo'], '')

    def test_valores_invalidos_se_rechazan_con_su_motivo(self):
        respuesta = self._importar_repuestos([
            {'codigo': 'X-1', 'nombre': 'A', 'iva_porcentaje_defecto': '12'},
            {'codigo': 'X-2', 'nombre': 'B', 'precio_venta': '-3'},
            {'codigo': 'X-3', 'nombre': 'C', 'categoria': 'JUGUETES'},
            {'codigo': '', 'nombre': 'Sin código'},
        ]).json()
        motivos = {e['fila']: e['motivo'] for e in respuesta['errores']}
        self.assertIn('IVA', motivos[2])
        self.assertIn('mayor o igual a 0', motivos[3])
        self.assertIn('Categoría', motivos[4])
        self.assertIn('código', motivos[5])
        self.assertFalse(Repuesto.objects.filter(empresa=self.empresa).exists())

    def test_codigo_repetido_en_el_archivo(self):
        respuesta = self._importar_repuestos([
            {'codigo': 'DUP-1', 'nombre': 'Primero'},
            {'codigo': 'dup-1', 'nombre': 'Segundo'},
        ]).json()
        self.assertEqual(respuesta['nuevos'], 1)
        self.assertIn('fila 2', respuesta['errores'][0]['motivo'])
        self.assertEqual(Repuesto.objects.get(empresa=self.empresa, codigo='DUP-1').nombre, 'Primero')

    def test_desactivado_no_se_reactiva(self):
        Repuesto.objects.create(empresa=self.empresa, codigo='OFF-1', nombre='Inactivo', is_active=False)
        respuesta = self._importar_repuestos([{'codigo': 'OFF-1', 'nombre': 'Inactivo'}]).json()
        self.assertIn('desactivado', respuesta['errores'][0]['motivo'])
        self.assertFalse(Repuesto.objects.get(empresa=self.empresa, codigo='OFF-1').is_active)

    def test_no_toca_el_mismo_codigo_de_otra_empresa(self):
        otra = Empresa.objects.create(nombre_comercial='Otra', ruc=str(uuid4().int % 10**13).zfill(13))
        Repuesto.objects.create(empresa=otra, codigo='COM-1', nombre='De la otra', precio_venta=Decimal('1'))
        respuesta = self._importar_repuestos([{'codigo': 'COM-1', 'nombre': 'Mío', 'precio_venta': '5'}]).json()
        self.assertEqual(respuesta['nuevos'], 1)
        self.assertEqual(Repuesto.objects.get(empresa=otra, codigo='COM-1').precio_venta, Decimal('1.00'))

    def test_vista_previa_no_guarda(self):
        self._importar_repuestos([{'codigo': 'PRE-1', 'nombre': 'Previa'}], validar=True)
        self.assertFalse(Repuesto.objects.filter(codigo='PRE-1').exists())

    def test_permisos(self):
        # CAJERO: inventario solo para ver. MECANICO: lo modifica.
        cajero = self._importar_repuestos([{'codigo': 'P-1', 'nombre': 'X'}], cliente=self._como('CAJERO'))
        self.assertEqual(cajero.status_code, 403)
        mecanico = self._importar_repuestos([{'codigo': 'P-1', 'nombre': 'X'}], cliente=self._como('MECANICO'))
        self.assertEqual(mecanico.status_code, 200, mecanico.content)

    def test_plantilla(self):
        libro = load_workbook(BytesIO(self.client.get(self.URL + 'plantilla/').content))
        self.assertEqual(libro.worksheets[0].max_row, 1)
        self.assertIn('Instrucciones', libro.sheetnames)


class ImportacionServiciosTests(_Base):
    URL = '/api/servicios/importar/'

    def test_crea_servicios_con_categoria_por_etiqueta(self):
        respuesta = self._importar(self.URL, ImportadorServicios, [
            {'codigo': 'mo-1', 'nombre': 'Cambio de aceite', 'categoria': 'Mantenimiento Preventivo',
             'tiempo_estimado_minutos': '45', 'precio_referencial': '15', 'iva_porcentaje_defecto': '0%'},
        ])
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        servicio = Servicio.objects.get(empresa=self.empresa, codigo='MO-1')
        self.assertEqual(servicio.categoria, Servicio.Categoria.MANTENIMIENTO)
        self.assertEqual(servicio.tiempo_estimado_minutos, 45)
        self.assertEqual(servicio.iva_porcentaje_defecto, Decimal('0.0000'))

    def test_tiempo_invalido(self):
        respuesta = self._importar(self.URL, ImportadorServicios, [
            {'codigo': 'MO-2', 'nombre': 'X', 'tiempo_estimado_minutos': 'una hora'},
        ]).json()
        self.assertIn('entero', respuesta['errores'][0]['motivo'])

    def test_plantilla(self):
        respuesta = self.client.get(self.URL + 'plantilla/')
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('spreadsheetml', respuesta['Content-Type'])
