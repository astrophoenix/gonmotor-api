"""El dueño actual de un vehículo es por empresa.

El vehículo es global (la placa es única en el sistema) y varias empresas
pueden atenderlo, pero cada cliente pertenece a una sola empresa. Por eso cada
empresa tiene su propio dueño actual: asignar el de una no cierra el de otra, y
ninguna lectura (ficha, listado, exportación, avisos al cliente) debe
mostrar o usar el cliente de otra empresa.
"""

import json
from io import BytesIO
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.authentication.models import UsuarioEmpresa
from apps.clientes.models import Cliente
from apps.core.testing import fijar_empresa
from apps.empresas.models import Empresa, Taller
from apps.notificaciones.models import NotificacionProgramada
from apps.vehiculos.models import Vehiculo, VehiculoPropietario, asignar_propietario_actual, vincular_empresa

User = get_user_model()


def _ruc():
    return str(uuid4().int % 10**13).zfill(13)


class PropietarioPorEmpresaTests(TestCase):

    def setUp(self):
        self.empresa_a = Empresa.objects.create(nombre_comercial='Taller A', ruc=_ruc())
        self.empresa_b = Empresa.objects.create(nombre_comercial='Taller B', ruc=_ruc())
        self.vehiculo = Vehiculo.objects.create(placa='PRO0001', marca='Kia', modelo='Rio')
        vincular_empresa(self.vehiculo, self.empresa_a.pk)
        vincular_empresa(self.vehiculo, self.empresa_b.pk)
        self.cliente_a = Cliente.objects.create(
            empresa=self.empresa_a, identificacion='1711111111', nombre='Cliente de A', telefono='0991111111',
        )
        self.cliente_b = Cliente.objects.create(
            empresa=self.empresa_b, identificacion='1722222222', nombre='Cliente de B', telefono='0992222222',
        )

    def _api(self, empresa):
        user = User.objects.create_user(username=f'admin-{uuid4().hex[:6]}', password='x')
        relacion = UsuarioEmpresa.objects.create(user=user, empresa=empresa, rol='ADMIN_EMPRESA')
        relacion.talleres.set([Taller.objects.create(empresa=empresa, nombre='T', codigo_sucursal=uuid4().hex[:3], direccion='x')])
        cliente = APIClient()
        cliente.force_authenticate(user=user)
        fijar_empresa(cliente, str(empresa.uuid))
        return cliente

    def _dos_duenos(self):
        """Cada empresa con su dueño actual (A primero: el orden importa para `.first()`)."""
        asignar_propietario_actual(self.vehiculo, self.cliente_a)
        asignar_propietario_actual(self.vehiculo, self.cliente_b)

    # --- Escritura -------------------------------------------------------------

    def test_asignar_el_dueno_de_una_empresa_no_cierra_el_de_otra(self):
        self._dos_duenos()
        actuales = set(VehiculoPropietario.objects.filter(vehiculo=self.vehiculo, es_actual=True)
                       .values_list('cliente_id', flat=True))
        self.assertEqual(actuales, {self.cliente_a.pk, self.cliente_b.pk})

    def test_cambiar_de_dueno_dentro_de_la_empresa_cierra_al_anterior(self):
        self._dos_duenos()
        otro_b = Cliente.objects.create(empresa=self.empresa_b, identificacion='1733333333', nombre='Otro de B')
        asignar_propietario_actual(self.vehiculo, otro_b)
        actuales = set(VehiculoPropietario.objects.filter(vehiculo=self.vehiculo, es_actual=True)
                       .values_list('cliente_id', flat=True))
        self.assertEqual(actuales, {self.cliente_a.pk, otro_b.pk})

    # --- Lectura ---------------------------------------------------------------

    def test_la_ficha_muestra_el_dueno_de_la_propia_empresa(self):
        self._dos_duenos()
        url = f'/api/vehiculos/{self.vehiculo.uuid}/'
        self.assertEqual(self._api(self.empresa_a).get(url).json()['cliente_nombre'], 'Cliente de A')
        self.assertEqual(self._api(self.empresa_b).get(url).json()['cliente_nombre'], 'Cliente de B')

    def test_el_listado_muestra_el_dueno_de_la_propia_empresa(self):
        self._dos_duenos()
        filas = self._api(self.empresa_b).get('/api/vehiculos/').json()['results']
        self.assertEqual([f['cliente_nombre'] for f in filas], ['Cliente de B'])

    def test_la_exportacion_no_trae_clientes_de_otra_empresa(self):
        self._dos_duenos()
        contenido = self._api(self.empresa_b).get('/api/vehiculos/export-excel/').content
        textos = [str(c.value) for fila in load_workbook(BytesIO(contenido)).active.iter_rows() for c in fila if c.value]
        self.assertFalse([t for t in textos if 'Cliente de A' in t])

    def test_el_filtro_por_cliente_no_acepta_clientes_de_otra_empresa(self):
        self._dos_duenos()
        filas = self._api(self.empresa_b).get('/api/vehiculos/', {'cliente': str(self.cliente_a.uuid)}).json()['results']
        self.assertEqual(filas, [])

    def test_el_mantenimiento_programado_usa_el_cliente_de_la_propia_empresa(self):
        self._dos_duenos()
        respuesta = self._api(self.empresa_b).post(
            '/api/notificaciones/programadas/',
            {'vehiculo_id': str(self.vehiculo.uuid), 'km_objetivo': 10000},
            format='json',
        )
        self.assertEqual(respuesta.status_code, 201, respuesta.data)
        notificacion = NotificacionProgramada.objects.get(uuid=respuesta.data['id'])
        self.assertEqual(notificacion.cliente, self.cliente_b)
        self.assertEqual(notificacion.empresa, self.empresa_b)

    def test_el_mantenimiento_no_se_programa_sin_dueno_en_la_empresa(self):
        asignar_propietario_actual(self.vehiculo, self.cliente_a)
        respuesta = self._api(self.empresa_b).post(
            '/api/notificaciones/programadas/',
            {'vehiculo_id': str(self.vehiculo.uuid), 'km_objetivo': 10000},
            format='json',
        )
        self.assertEqual(respuesta.status_code, 400)
        self.assertFalse(NotificacionProgramada.objects.exists())

    def test_agregar_el_vehiculo_desde_un_cliente_no_quita_el_dueno_de_otra_empresa(self):
        asignar_propietario_actual(self.vehiculo, self.cliente_a)
        respuesta = self._api(self.empresa_b).patch(
            f'/api/clientes/{self.cliente_b.uuid}/',
            # Como la web: multipart con los vehículos serializados en JSON.
            {'vehiculos': json.dumps([{'placa': 'PRO0001', 'marca': 'Kia', 'modelo': 'Rio'}])},
            format='multipart',
        )
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        actuales = set(VehiculoPropietario.objects.filter(vehiculo=self.vehiculo, es_actual=True)
                       .values_list('cliente_id', flat=True))
        self.assertEqual(actuales, {self.cliente_a.pk, self.cliente_b.pk})

    def test_agregar_el_vehiculo_desde_un_cliente_cierra_al_dueno_anterior_de_la_empresa(self):
        anterior_b = Cliente.objects.create(empresa=self.empresa_b, identificacion='1744444444', nombre='Anterior de B')
        asignar_propietario_actual(self.vehiculo, anterior_b)
        self._api(self.empresa_b).patch(
            f'/api/clientes/{self.cliente_b.uuid}/',
            {'vehiculos': json.dumps([{'placa': 'PRO0001', 'marca': 'Kia', 'modelo': 'Rio'}])},
            format='multipart',
        )
        actuales_b = list(VehiculoPropietario.objects.filter(
            vehiculo=self.vehiculo, es_actual=True, cliente__empresa=self.empresa_b,
        ).values_list('cliente_id', flat=True))
        self.assertEqual(actuales_b, [self.cliente_b.pk])
