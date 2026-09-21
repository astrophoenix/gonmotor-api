from django.contrib.auth.models import User
from django.test import TestCase
from rest_framework.test import APITestCase, APIClient
from openpyxl import load_workbook
from io import BytesIO
from apps.empresas.models import Empresa
from apps.clientes.models import Cliente
from apps.vehiculos.models import Vehiculo, VehiculoPropietario
from apps.vehiculos.serializers import VehiculoNestedSerializer
from apps.vehiculos.views import formatear_placa


class ClientePropietarioSerializerTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='tester', password='x')
        self.empresa = Empresa.objects.create(
            nombre_comercial='Taller GON',
            razon_social='GON SA',
            ruc='1750000000001',
        )
        self.cliente1 = Cliente.objects.create(
            empresa=self.empresa,
            tipo_identificacion='R',
            identificacion='0001',
            nombre='Cliente Uno',
        )
        self.cliente2 = Cliente.objects.create(
            empresa=self.empresa,
            tipo_identificacion='R',
            identificacion='0002',
            nombre='Cliente Dos',
        )

    def _create(self, cliente_id=None):
        data = {'placa': 'ABC123', 'marca': 'Toyota', 'modelo': 'Corolla', 'cliente_id': cliente_id}
        s = VehiculoNestedSerializer(data=data)
        self.assertTrue(s.is_valid(), s.errors)
        return s.save()

    def test_create_with_cliente(self):
        v = self._create(cliente_id=self.cliente1.id)
        rel = v.propietarios.get(es_actual=True)
        self.assertEqual(rel.cliente_id, self.cliente1.id)
        self.assertTrue(rel.es_actual)
        self.assertIsNone(rel.fecha_fin)

    def test_create_without_cliente(self):
        v = self._create(cliente_id=None)
        self.assertFalse(v.propietarios.filter(es_actual=True).exists())

    def test_update_switch_cliente_marks_previous_inactive(self):
        v = self._create(cliente_id=self.cliente1.id)
        s = VehiculoNestedSerializer(
            instance=v,
            data={'cliente_id': self.cliente2.id},
            partial=True,
        )
        self.assertTrue(s.is_valid(), s.errors)
        s.save()
        self.assertFalse(v.propietarios.filter(es_actual=True, cliente_id=self.cliente1.id).exists())
        self.assertTrue(v.propietarios.filter(es_actual=True, cliente_id=self.cliente2.id).exists())

    def test_update_clear_cliente(self):
        v = self._create(cliente_id=self.cliente1.id)
        s = VehiculoNestedSerializer(
            instance=v,
            data={'cliente_id': ''},
            partial=True,
        )
        self.assertTrue(s.is_valid(), s.errors)
        s.save()
        self.assertFalse(v.propietarios.filter(es_actual=True).exists())
        self.assertTrue(v.propietarios.filter(es_actual=False).exists())


class VehiculoViewSetFilterTest(APITestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username='tester', password='x')
        self.empresa = Empresa.objects.create(
            nombre_comercial='Taller GON',
            razon_social='GON SA',
            ruc='1750000000002',
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        self.client.defaults['HTTP_X_EMPRESA_ID'] = self.empresa.pk

        self.vehiculo = Vehiculo.objects.create(
            placa='GSU5812', marca='Toyota', modelo='Corolla', anio=2018,
        )
        self.vehiculo.empresas.add(self.empresa)

        self.vehiculo_otro = Vehiculo.objects.create(
            placa='PBA1234', marca='Chevrolet', modelo='Sail', anio=2020,
        )
        self.vehiculo_otro.empresas.add(self.empresa)

        self.vehiculo_inactivo = Vehiculo.objects.create(
            placa='XYZ9999', marca='Mazda', modelo='3', anio=2015, is_active=False,
        )
        self.vehiculo_inactivo.empresas.add(self.empresa)

    def test_search_placa_con_guion(self):
        resp = self.client.get('/api/vehiculos/', {'search': 'GSU-5812'})
        self.assertEqual(resp.status_code, 200)
        placas = [v['placa'] for v in resp.json()['results']]
        self.assertIn('GSU5812', placas)

    def test_search_placa_sin_guion(self):
        resp = self.client.get('/api/vehiculos/', {'search': 'GSU5812'})
        placas = [v['placa'] for v in resp.json()['results']]
        self.assertIn('GSU5812', placas)
        self.assertNotIn('PBA1234', placas)

    def test_filtro_anio(self):
        resp = self.client.get('/api/vehiculos/', {'anio': 2020})
        placas = [v['placa'] for v in resp.json()['results']]
        self.assertEqual(placas, ['PBA1234'])

    def test_listado_oculta_inactivos_por_defecto(self):
        resp = self.client.get('/api/vehiculos/')
        placas = [v['placa'] for v in resp.json()['results']]
        self.assertNotIn('XYZ9999', placas)

    def test_filtro_estado_inactivo(self):
        resp = self.client.get('/api/vehiculos/', {'estado': 'inactivo'})
        self.assertEqual(resp.status_code, 200)
        placas = [v['placa'] for v in resp.json()['results']]
        self.assertIn('XYZ9999', placas)

    def test_filtro_estado_activo(self):
        resp = self.client.get('/api/vehiculos/', {'estado': 'activo'})
        placas = [v['placa'] for v in resp.json()['results']]
        self.assertIn('GSU5812', placas)
        self.assertNotIn('XYZ9999', placas)


class VehiculoExportTest(APITestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username='tester', password='x')
        self.empresa = Empresa.objects.create(
            nombre_comercial='Taller GON',
            razon_social='GON SA',
            ruc='1750000000003',
        )
        self.cliente = Cliente.objects.create(
            empresa=self.empresa,
            tipo_identificacion='R',
            identificacion='0003',
            nombre='Propietario Demo',
        )
        self.vehiculo = Vehiculo.objects.create(
            placa='GSU5812', marca='Toyota', modelo='Corolla', anio=2018,
        )
        self.vehiculo.empresas.add(self.empresa)
        VehiculoPropietario.objects.create(
            vehiculo=self.vehiculo, cliente=self.cliente, es_actual=True,
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        self.client.defaults['HTTP_X_EMPRESA_ID'] = self.empresa.pk

    def test_exportar_pdf(self):
        resp = self.client.get('/api/vehiculos/exportar-pdf/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        contenido = b''.join(resp.streaming_content)
        self.assertTrue(contenido.startswith(b'%PDF'))

    def test_exportar_excel(self):
        resp = self.client.get('/api/vehiculos/export-excel/')
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.content.startswith(b'PK'))

    def test_formatear_placa(self):
        self.assertEqual(formatear_placa('GSU5812'), 'GSU-5812')
        self.assertEqual(formatear_placa('ABC123'), 'ABC-123')
        self.assertEqual(formatear_placa('GSU-5812'), 'GSU-5812')
        self.assertEqual(formatear_placa(''), '')
        self.assertEqual(formatear_placa(None), '')

    def test_excel_placa_formateada(self):
        resp = self.client.get('/api/vehiculos/export-excel/')
        wb = load_workbook(BytesIO(resp.content))
        valores = [
            celda
            for fila in wb.active.iter_rows(values_only=True)
            for celda in fila
        ]
        self.assertIn('GSU-5812', valores)
        self.assertNotIn('GSU5812', valores)

    def test_listado_marca_dueno_inactivo(self):
        self.cliente.is_active = False
        self.cliente.save(update_fields=['is_active'])
        resp = self.client.get('/api/vehiculos/')
        item = next(v for v in resp.json()['results'] if v['placa'] == 'GSU5812')
        self.assertEqual(item['cliente_nombre'], 'Propietario Demo (inactivo)')

    def test_listado_dueno_activo_sin_marca(self):
        resp = self.client.get('/api/vehiculos/')
        item = next(v for v in resp.json()['results'] if v['placa'] == 'GSU5812')
        self.assertEqual(item['cliente_nombre'], 'Propietario Demo')

    def test_excel_dueno_inactivo(self):
        self.cliente.is_active = False
        self.cliente.save(update_fields=['is_active'])
        resp = self.client.get('/api/vehiculos/export-excel/')
        wb = load_workbook(BytesIO(resp.content))
        valores = [
            celda
            for fila in wb.active.iter_rows(values_only=True)
            for celda in fila
        ]
        self.assertIn('Propietario Demo (inactivo)', valores)