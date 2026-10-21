from datetime import timedelta
from decimal import Decimal

from django.test import SimpleTestCase, TestCase
from django.utils import timezone

from .models import InspeccionVehiculo
from apps.core.testing import fijar_empresa, firma_data_url, imagen_de_prueba


class OrdenesCrudTests(SimpleTestCase):
    def test_imports_crud_components(self):
        from apps.ordenes.views import OrdenTrabajoViewSet
        from apps.ordenes.serializers import OrdenTrabajoSerializer

        self.assertIsNotNone(OrdenTrabajoViewSet)
        self.assertIsNotNone(OrdenTrabajoSerializer)


class InspeccionFechasTests(TestCase):
    def _inspeccion(self, **atributos):
        return InspeccionVehiculo.objects.create(
            numero_inspeccion=atributos.pop('numero_inspeccion', 'INS-TEST-0001'),
            **atributos,
        )

    def test_fecha_inspeccion_por_defecto(self):
        inspeccion = self._inspeccion()

        self.assertIsNotNone(inspeccion.fecha_inspeccion)

    def test_fecha_finalizacion_nula_mientras_no_este_finalizada(self):
        inspeccion = self._inspeccion(estado='EN_PROCESO')

        self.assertIsNone(inspeccion.fecha_finalizacion)

    def test_finalizar_sella_fecha_finalizacion(self):
        inspeccion = self._inspeccion(estado='EN_PROCESO')

        inspeccion.estado = 'FINALIZADA'
        inspeccion.save(update_fields=['estado', 'updated_at'])

        inspeccion.refresh_from_db()
        self.assertIsNotNone(inspeccion.fecha_finalizacion)
        self.assertEqual(inspeccion.estado, 'FINALIZADA')

    def test_finalizar_respeta_update_fields(self):
        inspeccion = self._inspeccion(estado='EN_PROCESO')

        inspeccion.estado = 'FINALIZADA'
        inspeccion.save(update_fields=['estado', 'updated_at'])

        # La fecha se escribe aunque no esté en update_fields.
        inspeccion.refresh_from_db()
        self.assertIsNotNone(inspeccion.fecha_finalizacion)

    def test_reabrir_limpia_fecha_finalizacion(self):
        inspeccion = self._inspeccion(estado='EN_PROCESO')
        inspeccion.estado = 'FINALIZADA'
        inspeccion.save()
        self.assertIsNotNone(inspeccion.fecha_finalizacion)

        inspeccion.estado = 'EN_PROCESO'
        inspeccion.save()

        inspeccion.refresh_from_db()
        self.assertIsNone(inspeccion.fecha_finalizacion)

    def test_no_sobrescribe_fecha_finalizacion_al_editar_una_finalizada(self):
        inspeccion = self._inspeccion(estado='EN_PROCESO')
        inspeccion.estado = 'FINALIZADA'
        inspeccion.save()
        cerrada = inspeccion.fecha_finalizacion

        inspeccion.diagnostico_tecnico = 'Ajuste posterior'
        inspeccion.save()

        inspeccion.refresh_from_db()
        self.assertEqual(inspeccion.fecha_finalizacion, cerrada)

    def test_duracion_inspeccion_en_minutos(self):
        inicio = timezone.now() - timedelta(hours=2, minutes=15)
        inspeccion = self._inspeccion(estado='EN_PROCESO', fecha_inspeccion=inicio)
        inspeccion.estado = 'FINALIZADA'
        inspeccion.save()

        self.assertEqual(inspeccion.duracion_inspeccion, 135)

    def test_duracion_calculada_entre_fechas(self):
        inicio = timezone.now() - timedelta(minutes=95)
        inspeccion = InspeccionVehiculo(
            numero_inspeccion='INS-TEST-0002',
            estado='FINALIZADA',
            fecha_inspeccion=inicio,
        )
        inspeccion.save()
        inspeccion.refresh_from_db()
        inspeccion.fecha_finalizacion = inicio + timedelta(minutes=95)
        inspeccion.save()

        self.assertEqual(inspeccion.duracion_inspeccion, 95)

    def test_duracion_es_none_sin_finalizacion(self):
        inspeccion = self._inspeccion(estado='EN_PROCESO')

        self.assertIsNone(inspeccion.duracion_inspeccion)


class InspeccionListadoFiltrosTests(TestCase):
    """El listado de inspecciones tiene panel de búsqueda: filtros por estado,
    tipo, taller y rango de fechas, y búsqueda por número, placa y cliente."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from rest_framework.test import APIClient

        from apps.clientes.models import Cliente
        from apps.empresas.models import Empresa, Taller
        from apps.vehiculos.models import Vehiculo
        from .models import RecepcionVehiculo

        self.empresa = Empresa.objects.create(nombre_comercial='Taller Filtros', ruc='666666666666')
        self.otra_empresa = Empresa.objects.create(nombre_comercial='Taller Ajeno', ruc='555555555555')
        self.taller = Taller.objects.create(
            empresa=self.empresa, nombre='Taller Norte', direccion='Av. Norte 100', prefijo_inspeccion='INS-'
        )
        self.cliente = Cliente.objects.create(
            empresa=self.empresa, nombre='Cliente Filtros', identificacion='0912345678'
        )
        self.vehiculo = Vehiculo.objects.create(placa='FIL-1234', marca='Toyota', modelo='Corolla')

        hoy = timezone.now()
        # Una inspección por recepción (restricción del modelo), así que cada
        # caso de prueba prepara la suya.
        self.recepcion_1 = self._recepcion('REC-FIL-0001')
        self.recepcion_2 = self._recepcion('REC-FIL-0002')

        self.pendiente = self._inspeccion(
            numero_inspeccion='INS-FIL-0001',
            recepcion=self.recepcion_1,
            estado='PENDIENTE',
            tipo_inspeccion='DIAGNOSTICO',
            fecha_inspeccion=hoy,
        )
        self.finalizada = self._inspeccion(
            numero_inspeccion='INS-FIL-0002',
            recepcion=self.recepcion_2,
            estado='FINALIZADA',
            tipo_inspeccion='MANTENIMIENTO',
            fecha_inspeccion=hoy - timedelta(days=10),
            fecha_finalizacion=hoy - timedelta(days=9),
        )
        # Inspección creada sin recepción: el vehículo y el cliente están en la
        # propia inspección y deben ser alcanzables por la búsqueda.
        self.cliente_standalone = Cliente.objects.create(
            empresa=self.empresa, nombre='Cliente Avenida', identificacion='0998765432'
        )
        self.vehiculo_standalone = Vehiculo.objects.create(placa='ABC4321', marca='Mazda', modelo='CX-5')
        self.standalone = self._inspeccion(
            numero_inspeccion='INS-FIL-0003',
            recepcion=None,
            cliente=self.cliente_standalone,
            vehiculo=self.vehiculo_standalone,
            sucursal=None,
            fecha_inspeccion=hoy,
        )
        self.de_otra_empresa = self._inspeccion(
            numero_inspeccion='INS-FIL-0004',
            empresa=self.otra_empresa,
            recepcion=self._recepcion('REC-FIL-0009', empresa=self.otra_empresa),
            cliente=Cliente.objects.create(empresa=self.otra_empresa, nombre='Cliente Ajeno'),
            vehiculo=Vehiculo.objects.create(placa='ZZZ9999'),
            fecha_inspeccion=hoy,
        )

        self.usuario = get_user_model().objects.create_superuser(
            username='filtros', email='filtros@test.local', password='x'
        )
        self.client = APIClient()
        self.client.force_authenticate(self.usuario)

    def _recepcion(self, numero, empresa=None):
        from .models import RecepcionVehiculo

        return RecepcionVehiculo.objects.create(
            empresa=empresa or self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_recepcion=numero,
        )

    def _inspeccion(self, empresa=None, recepcion=None, cliente=None, vehiculo=None, sucursal='__default__', **atributos):
        if sucursal == '__default__':
            sucursal = self.taller if recepcion is not None else None
        return InspeccionVehiculo.objects.create(
            empresa=empresa or self.empresa,
            sucursal=sucursal,
            cliente=cliente or self.cliente,
            vehiculo=vehiculo or self.vehiculo,
            recepcion=recepcion,
            orden_trabajo=None,
            motivo_ingreso='Ruido en el motor',
            numero_inspeccion=atributos.pop('numero_inspeccion', 'INS-FIL-9999'),
            **atributos,
        )

    def _listar(self, **params):
        respuesta = self.client.get('/api/ordenes/inspecciones/', params)
        self.assertEqual(respuesta.status_code, 200)
        return respuesta.json()

    def _numeros(self, **params):
        return sorted(item['numero_inspeccion'] for item in self._listar(**params)['results'])

    def test_no_muestra_inspecciones_de_otra_empresa(self):
        self.assertEqual(self._listar()['count'], 3)

    def test_filtra_por_estado(self):
        self.assertEqual(self._numeros(estado='PENDIENTE'), ['INS-FIL-0001', 'INS-FIL-0003'])
        self.assertEqual(self._numeros(estado='FINALIZADA'), ['INS-FIL-0002'])

    def test_filtra_por_tipo(self):
        self.assertEqual(self._numeros(tipo_inspeccion='MANTENIMIENTO'), ['INS-FIL-0002'])

    def test_filtra_por_taller(self):
        self.assertEqual(self._numeros(sucursal=str(self.taller.uuid)), ['INS-FIL-0001', 'INS-FIL-0002'])
        self.assertEqual(self._listar(sucursal=999)['count'], 0)

    def test_filtra_por_rango_de_fechas(self):
        # La fecha local del taller, no la de UTC: entre las 00:00 y las 05:00
        # UTC el filtro por fecha_inspeccion__date no ve el mismo día.
        hoy = timezone.localdate().isoformat()
        ayer = (timezone.localdate() - timedelta(days=1)).isoformat()

        self.assertEqual(self._numeros(fecha_desde=hoy), ['INS-FIL-0001', 'INS-FIL-0003'])
        self.assertEqual(self._numeros(fecha_hasta=ayer), ['INS-FIL-0002'])
        self.assertEqual(
            self._numeros(fecha_desde=ayer, fecha_hasta=hoy),
            ['INS-FIL-0001', 'INS-FIL-0003'],
        )

    def test_filtra_inspecciones_sin_recepcion(self):
        self.assertEqual(self._numeros(sin_recepcion='1'), ['INS-FIL-0003'])

    def test_ignora_fechas_invalidas(self):
        self.assertEqual(self._listar(fecha_desde='no-es-fecha')['count'], 3)

    def test_combina_filtros(self):
        self.assertEqual(
            self._numeros(estado='FINALIZADA', tipo_inspeccion='MANTENIMIENTO', sucursal=str(self.taller.uuid)),
            ['INS-FIL-0002'],
        )

    def test_busca_por_numero_placa_cliente_y_recepcion(self):
        casos = {
            'INS-FIL-0001': ['INS-FIL-0001'],
            'FIL-1234': ['INS-FIL-0001', 'INS-FIL-0002'],
            'ABC4321': ['INS-FIL-0003'],
            'Cliente Filtros': ['INS-FIL-0001', 'INS-FIL-0002'],
            'Cliente Avenida': ['INS-FIL-0003'],
            'REC-FIL-0001': ['INS-FIL-0001'],
        }
        for termino, esperados in casos.items():
            with self.subTest(termino=termino):
                self.assertEqual(self._numeros(search=termino), esperados)

    def test_expone_la_orden_de_trabajo_en_la_relacion(self):
        from .models import OrdenTrabajo

        orden = OrdenTrabajo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_orden='OT-FIL-0001',
            estado='EN_PROCESO',
            fecha_ingreso=timezone.now(),
        )
        self.finalizada.orden_trabajo = orden
        self.finalizada.save(update_fields=['orden_trabajo'])

        item = next(i for i in self._listar()['results'] if i['numero_inspeccion'] == 'INS-FIL-0002')

        self.assertEqual(item['orden_trabajo'], str(orden.uuid))
        self.assertEqual(item['orden_trabajo_numero'], 'OT-FIL-0001')
        self.assertEqual(item['orden_trabajo_estado'], 'EN_PROCESO')
        self.assertTrue(item['orden_trabajo_estado_display'])
        self.assertTrue(item['tiene_orden_trabajo'])

    def test_inspeccion_sin_recepcion_no_inventa_relaciones(self):
        item = next(i for i in self._listar()['results'] if i['numero_inspeccion'] == 'INS-FIL-0003')

        self.assertIsNone(item['recepcion'])
        self.assertIsNone(item['orden_trabajo'])
        self.assertIsNone(item['orden_trabajo_numero'])
        self.assertIsNone(item['cotizacion_numero'])
        self.assertEqual(item['vehiculo']['placa'], 'ABC4321')
        self.assertEqual(item['cliente']['nombre'], 'Cliente Avenida')

    def test_expone_el_kilometraje_del_vehiculo(self):
        self.vehiculo.kilometraje_actual = 72500
        self.vehiculo.save(update_fields=['kilometraje_actual'])
        self.vehiculo_standalone.kilometraje_actual = 12300
        self.vehiculo_standalone.save(update_fields=['kilometraje_actual'])

        resultados = {i['numero_inspeccion']: i for i in self._listar()['results']}

        # Con recepción: el vehículo de la propia inspección y el anidado.
        con_recepcion = resultados['INS-FIL-0001']
        self.assertEqual(con_recepcion['vehiculo']['kilometraje_actual'], 72500)
        self.assertEqual(con_recepcion['recepcion']['vehiculo']['kilometraje_actual'], 72500)
        # Sin recepción: solo el vehículo de la inspección.
        self.assertEqual(resultados['INS-FIL-0003']['vehiculo']['kilometraje_actual'], 12300)


class InspeccionKilometrajeDiagnosticoTests(TestCase):
    """La inspección guarda su propia lectura del odómetro porque el vehículo
    puede salir a prueba de ruta. Al finalizar la sube al vehículo, y el
    odómetro maestro nunca retrocede."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from rest_framework.test import APIClient

        from apps.clientes.models import Cliente
        from apps.empresas.models import Empresa, Taller
        from apps.vehiculos.models import Vehiculo
        from .models import RecepcionVehiculo

        self.empresa = Empresa.objects.create(nombre_comercial='Taller Km', ruc='777777777777')
        self.taller = Taller.objects.create(
            empresa=self.empresa, nombre='Taller Centro', direccion='Av. Central 1', prefijo_inspeccion='KM-'
        )
        self.cliente = Cliente.objects.create(
            empresa=self.empresa, nombre='Cliente Km', identificacion='0911111111'
        )
        self.vehiculo = Vehiculo.objects.create(
            placa='KM-1234', marca='Toyota', modelo='Corolla', kilometraje_actual=50000
        )
        self.recepcion = RecepcionVehiculo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_recepcion='REC-KM-0001',
            kilometraje_ingreso=50000,
            estado='ACEPTADA',
        )
        self.inspeccion = self._inspeccion()

        self.usuario = get_user_model().objects.create_superuser(
            username='km', email='km@test.local', password='x'
        )
        self.client = APIClient()
        self.client.force_authenticate(self.usuario)

    def _inspeccion(self, **atributos):
        from .models import InspeccionVehiculo

        atributos = {
            'empresa': self.empresa,
            'sucursal': self.taller,
            'cliente': self.cliente,
            'vehiculo': self.vehiculo,
            'recepcion': self.recepcion,
            'motivo_ingreso': 'Ruido en el motor',
            **atributos,
        }
        if atributos['recepcion'] is None:
            atributos['sucursal'] = None
        return InspeccionVehiculo.objects.create(**atributos)

    def _actualizar(self, payload):
        return self.client.patch(
            f'/api/ordenes/inspecciones/{self.inspeccion.uuid}/',
            data=payload,
            format='json',
        )

    def test_precarga_el_odometro_del_vehiculo_al_crear(self):
        self.assertEqual(self.inspeccion.kilometraje_diagnostico, 50000)

    def test_al_finalizar_sube_el_odometro_del_vehiculo(self):
        respuesta = self._actualizar({
            'kilometraje_diagnostico': 52400,
            'estado': 'EN_PROCESO',
        })
        self.assertEqual(respuesta.status_code, 200)
        respuesta = self._actualizar({
            'kilometraje_diagnostico': 52400,
            'estado': 'FINALIZADA',
        })
        self.assertEqual(respuesta.status_code, 200)

        self.vehiculo.refresh_from_db()
        self.assertEqual(self.vehiculo.kilometraje_actual, 52400)

    def test_una_lectura_menor_no_hace_retroceder_el_vehiculo(self):
        self.vehiculo.kilometraje_actual = 60000
        self.vehiculo.save(update_fields=['kilometraje_actual'])
        inspeccion = self._inspeccion(
            numero_inspeccion='KM-0002', recepcion=None, kilometraje_diagnostico=50000, estado='PENDIENTE'
        )

        inspeccion.kilometraje_diagnostico = 50000
        inspeccion.estado = 'FINALIZADA'
        inspeccion.save()

        self.vehiculo.refresh_from_db()
        self.assertEqual(self.vehiculo.kilometraje_actual, 60000)

    def test_la_api_rechaza_un_kilometraje_menor_al_del_vehiculo(self):
        respuesta = self._actualizar({'kilometraje_diagnostico': 40000})

        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('kilometraje_diagnostico', respuesta.json())

    def test_la_api_permite_conservar_la_lectura_ya_guardada(self):
        # El vehículo puede haber subido en otra visita: la inspección conserva
        # su propia lectura sin que el guardado sea rechazado.
        self.vehiculo.kilometraje_actual = 58000
        self.vehiculo.save(update_fields=['kilometraje_actual'])

        respuesta = self._actualizar({'motivo_ingreso': 'Se reintervisita la falla'})

        self.assertEqual(respuesta.status_code, 200)
        self.inspeccion.refresh_from_db()
        self.assertEqual(self.inspeccion.kilometraje_diagnostico, 50000)

    def test_el_listado_expone_la_lectura_de_la_inspeccion(self):
        self.inspeccion.kilometraje_diagnostico = 52310
        self.inspeccion.save(update_fields=['kilometraje_diagnostico'])

        respuesta = self.client.get('/api/ordenes/inspecciones/')

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.json()['results'][0]['kilometraje_diagnostico'], 52310)

    def test_la_recepcion_no_hace_retroceder_el_odometro(self):
        self.vehiculo.kilometraje_actual = 70000
        self.vehiculo.save(update_fields=['kilometraje_actual'])
        self.recepcion.estado = 'PENDIENTE'
        self.recepcion.kilometraje_ingreso = 45000
        self.recepcion.save()

        # Se revalida el flujo de la vista, que es quien sincroniza.
        from apps.ordenes.views import RecepcionVehiculoViewSet

        RecepcionVehiculoViewSet()._sincronizar_kilometraje_vehiculo(self.recepcion)

        self.vehiculo.refresh_from_db()
        self.assertEqual(self.vehiculo.kilometraje_actual, 70000)


class OrdenTrabajoEstadosTests(TestCase):
    """Estados de la orden de trabajo: pendiente de inicio, en espera, en
    proceso, completado, entregado y cancelado, con cambios libres entre
    estados (las únicas reglas obligatorias son el motivo en espera y la
    fecha de entrega al entregar)."""

    def setUp(self):
        from apps.clientes.models import Cliente
        from apps.empresas.models import Empresa
        from apps.vehiculos.models import Vehiculo

        self.empresa = Empresa.objects.create(nombre_comercial='Taller Estados', ruc='888888888888')
        self.cliente = Cliente.objects.create(
            empresa=self.empresa, nombre='Cliente Estados', identificacion='0912345679'
        )
        self.vehiculo = Vehiculo.objects.create(placa='EST-0001', marca='Kia', modelo='Sportage')
        self.vehiculo.empresas.add(self.empresa)
        self._contador = 0

    def _orden(self, estado='PENDIENTE', motivo_espera=None):
        from .models import OrdenTrabajo

        self._contador += 1
        return OrdenTrabajo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_orden=f'OT-EST-{self._contador:04d}',
            estado=estado,
            motivo_espera=motivo_espera,
        )

    def _serializer(self, orden, payload):
        from .serializers import OrdenTrabajoSerializer

        return OrdenTrabajoSerializer(orden, data=payload, partial=True)

    def test_estado_por_defecto_pendiente(self):
        from .models import OrdenTrabajo

        self._contador += 1
        orden = OrdenTrabajo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_orden=f'OT-EST-{self._contador:04d}',
        )

        self.assertEqual(orden.estado, 'PENDIENTE')

    def test_espera_requiere_motivo(self):
        orden = self._orden(estado='EN_PROCESO')

        serializer = self._serializer(orden, {'estado': 'EN_ESPERA'})

        self.assertFalse(serializer.is_valid())
        self.assertIn('motivo_espera', serializer.errors)

    def test_espera_con_motivo_es_valida(self):
        orden = self._orden(estado='EN_PROCESO')

        serializer = self._serializer(orden, {
            'estado': 'EN_ESPERA',
            'motivo_espera': 'Esperando repuestos de frenos',
        })

        self.assertTrue(serializer.is_valid(), serializer.errors)

    def test_se_puede_saltar_etapas(self):
        orden = self._orden(estado='PENDIENTE')

        serializer = self._serializer(orden, {'estado': 'ENTREGADO'})

        self.assertTrue(serializer.is_valid(), serializer.errors)
        guardada = serializer.save()
        self.assertEqual(guardada.estado, 'ENTREGADO')

    def test_espera_puede_volver_a_proceso_o_cancelarse(self):
        orden = self._orden(estado='EN_ESPERA', motivo_espera='Esperando repuestos')
        serializer = self._serializer(orden, {'estado': 'CANCELADO'})
        self.assertTrue(serializer.is_valid(), serializer.errors)
        orden2 = self._orden(estado='EN_ESPERA', motivo_espera='Esperando aprobación')
        serializer2 = self._serializer(orden2, {'estado': 'EN_PROCESO'})
        self.assertTrue(serializer2.is_valid(), serializer2.errors)

    def test_entregar_sella_fecha_entrega(self):
        orden = self._orden(estado='COMPLETADO')
        serializer = self._serializer(orden, {'estado': 'ENTREGADO'})
        self.assertTrue(serializer.is_valid(), serializer.errors)
        guardada = serializer.save()
        self.assertIsNotNone(guardada.fecha_entrega)
        self.assertEqual(guardada.estado, 'ENTREGADO')

    def test_entregado_y_cancelado_pueden_reabrirse(self):
        entregada = self._orden(estado='ENTREGADO')
        serializer = self._serializer(entregada, {'estado': 'EN_PROCESO'})
        self.assertTrue(serializer.is_valid(), serializer.errors)
        self.assertEqual(serializer.validated_data['estado'], 'EN_PROCESO')

        cancelada = self._orden(estado='CANCELADO')
        serializer2 = self._serializer(cancelada, {'estado': 'PENDIENTE'})
        self.assertTrue(serializer2.is_valid(), serializer2.errors)
        self.assertEqual(serializer2.validated_data['estado'], 'PENDIENTE')

    def test_completado_puede_reaperturarse_a_proceso(self):
        orden = self._orden(estado='COMPLETADO')
        serializer = self._serializer(orden, {'estado': 'EN_PROCESO'})
        self.assertTrue(serializer.is_valid(), serializer.errors)
        self.assertEqual(serializer.validated_data['estado'], 'EN_PROCESO')

    def test_api_crea_orden_independiente_sin_cotizacion(self):
        """Una OT independiente se crea por POST sin referencia de cotización."""
        from django.contrib.auth import get_user_model
        from rest_framework.test import APIClient

        from apps.empresas.models import Taller
        from .models import OrdenTrabajo

        Taller.objects.create(empresa=self.empresa, nombre='Taller OT', prefijo_ot='OT-')
        usuario = get_user_model().objects.create_superuser(
            username='crea-ot', email='crea-ot@test.local', password='x'
        )
        client = APIClient()
        client.force_authenticate(usuario)
        fijar_empresa(client, self.empresa)
        response = client.post(
            '/api/ordenes/ordenes-trabajo/',
            {
                'cliente': str(self.cliente.uuid),
                'vehiculo': str(self.vehiculo.uuid),
                'estado': 'PENDIENTE',
                'fecha_ingreso': '2026-10-01T08:30:00Z',
                'prioridad': 'MEDIA',
                'tipo_trabajo': 'MANTENIMIENTO',
                'kilometraje': self.vehiculo.kilometraje_actual,
                'mecanico_principal': str(usuario.profile.uuid),
            },
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.data)
        creada = OrdenTrabajo.objects.get(uuid=response.data['id'])
        self.assertEqual(creada.empresa_id, self.empresa.id)
        self.assertTrue(creada.numero_orden)
        self.assertIsNone(creada.cotizacion_origen_id)
        self.assertEqual(creada.estado, 'PENDIENTE')


class RecepcionRelacionesFlujoApiTests(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model
        from rest_framework.test import APIClient

        from apps.citas.models import Cita
        from apps.clientes.models import Cliente
        from apps.cotizaciones.models import Cotizacion
        from apps.empresas.models import Empresa
        from apps.vehiculos.models import Vehiculo
        from .models import OrdenTrabajo, RecepcionVehiculo

        self.empresa = Empresa.objects.create(nombre_comercial='Taller Relaciones', ruc='777777777777')
        self.cliente = Cliente.objects.create(empresa=self.empresa, nombre='Cliente Relaciones')
        self.vehiculo = Vehiculo.objects.create(placa='REL-0001', marca='Toyota', modelo='Corolla')
        self.vehiculo.empresas.add(self.empresa)
        self.recepcion = RecepcionVehiculo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_recepcion='REC-REL-0001',
        )
        self.cita = Cita.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            fecha_cita=timezone.localdate() + timedelta(days=1),
            hora_cita='10:00',
        )
        self.cotizacion = Cotizacion.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_cotizacion='COT-REL-0001',
        )
        self.orden = OrdenTrabajo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_orden='OT-REL-0001',
        )
        user = get_user_model().objects.create_superuser(
            username='relaciones', email='relaciones@test.local', password='x'
        )
        self.client = APIClient()
        self.client.force_authenticate(user=user)
        fijar_empresa(self.client, str(self.empresa.uuid))
    def _actualizar(self, tipo, entidad_id, accion):
        return self.client.post(
            f'/api/recepciones/{self.recepcion.uuid}/relaciones/',
            {'tipo': tipo, 'id': entidad_id, 'accion': accion},
            format='json',
        )

    def test_get_devuelve_relaciones_y_capacidades(self):
        self._actualizar('cita', str(self.cita.uuid), 'vincular')
        self.cotizacion.recepcion_origen = self.recepcion
        self.cotizacion.save(update_fields=['recepcion_origen'])
        self.recepcion.orden_trabajo = self.orden
        self.recepcion.save(update_fields=['orden_trabajo'])
        respuesta = self.client.get(f'/api/recepciones/{self.recepcion.uuid}/relaciones/')
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        datos = respuesta.json()
        for tipo, objeto in [('cita', self.cita), ('cotizacion', self.cotizacion), ('orden', self.orden)]:
            self.assertEqual(datos['relaciones'][tipo][0]['id'], str(objeto.uuid))
        self.assertFalse(any(datos['puede_agregar'].values()))

    def test_vincula_y_desvincula_cita_sin_borrarla(self):
        vincular = self._actualizar('cita', str(self.cita.uuid), 'vincular')
        self.assertEqual(vincular.status_code, 200, vincular.content)
        self.cita.refresh_from_db()
        self.assertEqual(self.cita.recepcion_generada_id, self.recepcion.pk)
        self.assertEqual(self.cita.estado, 'PROGRAMADA')

        desvincular = self._actualizar('cita', str(self.cita.uuid), 'desvincular')
        self.assertEqual(desvincular.status_code, 200, desvincular.content)
        self.cita.refresh_from_db()
        self.assertIsNone(self.cita.recepcion_generada_id)

    def test_vincula_y_desvincula_cotizacion_sin_borrarla(self):
        url = f'/api/cotizaciones/{self.cotizacion.uuid}/relaciones/'
        for accion, esperado in [('vincular', self.recepcion.pk), ('desvincular', None)]:
            respuesta = self.client.post(url, {'tipo': 'recepcion', 'id': str(self.recepcion.uuid), 'accion': accion}, format='json')
            self.assertEqual(respuesta.status_code, 200, respuesta.content)
            self.cotizacion.refresh_from_db()
            self.assertEqual(self.cotizacion.recepcion_origen_id, esperado)

    def test_vincula_varias_cotizaciones_activas_a_la_misma_recepcion(self):
        from apps.cotizaciones.models import Cotizacion
        segunda = Cotizacion.objects.create(empresa=self.empresa, cliente=self.cliente, vehiculo=self.vehiculo, numero_cotizacion='COT-REL-0002')
        for cotizacion in (self.cotizacion, segunda):
            respuesta = self.client.post(f'/api/cotizaciones/{cotizacion.uuid}/relaciones/', {'tipo': 'recepcion', 'id': str(self.recepcion.uuid), 'accion': 'vincular'}, format='json')
            self.assertEqual(respuesta.status_code, 200, respuesta.content)
        relaciones = self.client.get(f'/api/recepciones/{self.recepcion.uuid}/relaciones/').json()
        self.assertEqual({item['id'] for item in relaciones['relaciones']['cotizacion']}, {str(self.cotizacion.uuid), str(segunda.uuid)})
        self.assertFalse(relaciones['puede_agregar']['cotizacion'])

    def test_no_desvincula_cotizacion_indirecta_desde_recepcion(self):
        from apps.cotizaciones.models import Cotizacion

        inspeccion = InspeccionVehiculo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            recepcion=self.recepcion,
            motivo_ingreso='Inspección relacionada',
            numero_inspeccion='INS-REL-0003',
        )
        cotizacion = Cotizacion.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            inspeccion_origen=inspeccion,
            numero_cotizacion='COT-REL-0002',
        )

        relaciones = self.client.get(f'/api/recepciones/{self.recepcion.uuid}/relaciones/').json()
        self.assertIn(str(cotizacion.uuid), [item['id'] for item in relaciones['relaciones']['cotizacion']])
        respuesta = self._actualizar('cotizacion', str(cotizacion.uuid), 'desvincular')

        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        cotizacion.refresh_from_db()
        inspeccion.refresh_from_db()
        self.assertEqual(cotizacion.inspeccion_origen_id, inspeccion.pk)
        self.assertIsNone(cotizacion.recepcion_origen_id)
        self.assertEqual(inspeccion.recepcion_id, self.recepcion.pk)

    def test_no_edita_orden_desde_recepcion(self):
        for accion in ('vincular', 'desvincular'):
            respuesta = self._actualizar('orden', str(self.orden.uuid), accion)
            self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.recepcion.refresh_from_db()
        self.assertIsNone(self.recepcion.orden_trabajo_id)

    def test_no_permite_dos_inspecciones_en_la_misma_recepcion(self):
        primera = InspeccionVehiculo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            motivo_ingreso='Primera inspección',
            numero_inspeccion='INS-REL-0001',
        )
        segunda = InspeccionVehiculo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            motivo_ingreso='Segunda inspección',
            numero_inspeccion='INS-REL-0002',
        )
        self.assertEqual(self.client.post(f'/api/ordenes/inspecciones/{primera.uuid}/relaciones/', {'tipo': 'recepcion', 'id': str(self.recepcion.uuid), 'accion': 'vincular'}, format='json').status_code, 200)
        respuesta = self.client.post(f'/api/ordenes/inspecciones/{segunda.uuid}/relaciones/', {'tipo': 'recepcion', 'id': str(self.recepcion.uuid), 'accion': 'vincular'}, format='json')
        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        segunda.refresh_from_db()
        self.assertIsNone(segunda.recepcion_id)


class _BaseRelacionesFlujoApiTests(TestCase):
    """Base con el escenario mínimo del flujo taller: empresa, cliente, vehículo."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from rest_framework.test import APIClient

        from apps.citas.models import Cita
        from apps.clientes.models import Cliente
        from apps.cotizaciones.models import Cotizacion
        from apps.empresas.models import Empresa
        from apps.vehiculos.models import Vehiculo

        from .models import InspeccionVehiculo, OrdenTrabajo, RecepcionVehiculo

        self.Cita = Cita
        self.InspeccionVehiculo = InspeccionVehiculo
        self.OrdenTrabajo = OrdenTrabajo

        self.empresa = Empresa.objects.create(nombre_comercial='Taller Grafo', ruc='888888888888')
        self.cliente = Cliente.objects.create(empresa=self.empresa, nombre='Cliente Grafo')
        self.vehiculo = Vehiculo.objects.create(placa='GRA-0001', marca='Toyota', modelo='Corolla')
        self.vehiculo.empresas.add(self.empresa)
        self.recepcion = RecepcionVehiculo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_recepcion='REC-GRA-0001',
        )
        self.inspeccion = InspeccionVehiculo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_inspeccion='INS-GRA-0001',
        )
        self.cotizacion = Cotizacion.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_cotizacion='COT-GRA-0001',
        )
        self.orden = OrdenTrabajo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_orden='OT-GRA-0001',
        )
        user = get_user_model().objects.create_superuser(
            username='grafo', email='grafo@test.local', password='x'
        )
        self.client = APIClient()
        self.client.force_authenticate(user=user)
        fijar_empresa(self.client, str(self.empresa.uuid))
    def _post(self, url, tipo, entidad_id, accion='vincular'):
        return self.client.post(
            url, {'tipo': tipo, 'id': entidad_id, 'accion': accion}, format='json'
        )


class ListadoPdfFlujoApiTests(_BaseRelacionesFlujoApiTests):
    """``GET .../exportar-pdf/`` (listado) de inspecciones, cotizaciones y órdenes."""

    URLS = {
        'inspecciones': ('/api/ordenes/inspecciones/exportar-pdf/', 'listado_inspecciones.pdf'),
        'cotizaciones': ('/api/cotizaciones/exportar-pdf/', 'listado_cotizaciones.pdf'),
        'ordenes': ('/api/ordenes/ordenes-trabajo/exportar-pdf/', 'listado_ordenes.pdf'),
    }

    def _assert_pdf(self, respuesta, filename):
        self.assertEqual(respuesta.status_code, 200, getattr(respuesta, 'content', b''))
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertIn(filename, respuesta['Content-Disposition'])
        self.assertTrue(b''.join(respuesta.streaming_content).startswith(b'%PDF'))

    def test_descarga_listado_pdf(self):
        for modulo, (url, filename) in self.URLS.items():
            with self.subTest(modulo=modulo):
                self._assert_pdf(self.client.get(url), filename)

    def test_texto_con_caracteres_de_markup_no_rompe_el_pdf(self):
        # `&` y `<` sin escapar hacen fallar el Paragraph de reportlab.
        self.cliente.nombre = 'Pérez & Hijos <SA>'
        self.cliente.save()
        for modulo, (url, filename) in self.URLS.items():
            with self.subTest(modulo=modulo):
                self._assert_pdf(self.client.get(url), filename)

    def test_aplica_filtros_del_listado(self):
        for modulo, (url, filename) in self.URLS.items():
            with self.subTest(modulo=modulo):
                self._assert_pdf(self.client.get(url, {'search': 'NO-EXISTE', 'estado': 'PENDIENTE'}), filename)

    def test_requiere_autenticacion(self):
        from rest_framework.test import APIClient

        anonimo = APIClient()
        for modulo, (url, _) in self.URLS.items():
            with self.subTest(modulo=modulo):
                self.assertEqual(anonimo.get(url).status_code, 401)

    def test_el_pdf_de_una_cotizacion_sigue_disponible(self):
        respuesta = self.client.get(f'/api/cotizaciones/{self.cotizacion.uuid}/exportar-pdf/')
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('COT-GRA-0001', respuesta['Content-Disposition'])


class RecepcionListadoExportApiTests(_BaseRelacionesFlujoApiTests):
    """Exportes del listado de recepciones: empresa en sesión (no la de la OT),
    búsqueda y filtros del panel."""

    def _valores_excel(self, params=None):
        from io import BytesIO

        from openpyxl import load_workbook

        respuesta = self.client.get('/api/recepciones/export-excel/', params or {})
        self.assertEqual(respuesta.status_code, 200)
        contenido = b''.join(respuesta.streaming_content) if respuesta.streaming else respuesta.content
        hoja = load_workbook(BytesIO(contenido)).active
        return {str(c.value) for fila in hoja.iter_rows() for c in fila if c.value is not None}

    def test_pdf_incluye_recepciones_sin_orden(self):
        self.assertIsNone(self.recepcion.orden_trabajo_id)
        respuesta = self.client.get('/api/recepciones/exportar-pdf/')
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('listado_recepciones.pdf', respuesta['Content-Disposition'])
        self.assertTrue(b''.join(respuesta.streaming_content).startswith(b'%PDF'))

    def test_excel_incluye_recepciones_sin_orden(self):
        valores = self._valores_excel()
        self.assertIn('REC-GRA-0001', valores)
        self.assertIn('GRA-0001', valores)
        self.assertIn('Cliente Grafo', valores)

    def test_aplica_busqueda_y_filtros(self):
        self.assertNotIn('REC-GRA-0001', self._valores_excel({'search': 'NO-EXISTE'}))
        self.assertIn('REC-GRA-0001', self._valores_excel({'search': 'GRA-0001'}))
        self.assertNotIn('REC-GRA-0001', self._valores_excel({'solo_grua': 'true'}))

    def test_excluye_recepciones_de_otra_empresa(self):
        from apps.empresas.models import Empresa

        from .models import RecepcionVehiculo

        otra = Empresa.objects.create(nombre_comercial='Otra', ruc='777777777777')
        RecepcionVehiculo.objects.create(
            empresa=otra, cliente=self.cliente, vehiculo=self.vehiculo, numero_recepcion='REC-AJENA-1'
        )
        self.assertNotIn('REC-AJENA-1', self._valores_excel())


class ListadoExportLimiteApiTests(_BaseRelacionesFlujoApiTests):
    """Sobre el máximo de filas los exportes del flujo responden 400, no 500."""

    URLS = [
        '/api/ordenes/inspecciones/exportar-pdf/',
        '/api/cotizaciones/exportar-pdf/',
        '/api/ordenes/ordenes-trabajo/exportar-pdf/',
        '/api/recepciones/exportar-pdf/',
        '/api/recepciones/export-excel/',
    ]

    def test_sobre_el_limite_responde_400(self):
        from django.test import override_settings

        with override_settings(EXPORT_MAX_FILAS_PDF=0, EXPORT_MAX_FILAS_EXCEL=0):
            for url in self.URLS:
                with self.subTest(url=url):
                    respuesta = self.client.get(url)
                    self.assertEqual(respuesta.status_code, 400, respuesta.content)
                    self.assertIn('máximo', respuesta.json()['detail'])


class InspeccionRelacionesFlujoApiTests(_BaseRelacionesFlujoApiTests):
    def _url(self, inspeccion=None):
        return f'/api/ordenes/inspecciones/{(inspeccion or self.inspeccion).uuid}/relaciones/'

    def test_get_expone_recepcion_cotizacion_y_orden(self):
        respuesta = self.client.get(self._url())
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        datos = respuesta.json()
        self.assertEqual(set(datos['relaciones']), {'cita', 'recepcion', 'cotizacion', 'orden'})
        self.assertTrue(datos['puede_agregar']['recepcion'])

    def test_vincula_recepcion_y_cotizacion(self):
        respuesta = self._post(self._url(), 'recepcion', str(self.recepcion.uuid))
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        self.inspeccion.refresh_from_db()
        self.assertEqual(self.inspeccion.recepcion_id, self.recepcion.pk)
        respuesta = self._post(f'/api/cotizaciones/{self.cotizacion.uuid}/relaciones/', 'inspeccion', str(self.inspeccion.uuid))
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        self.cotizacion.refresh_from_db()
        self.assertEqual(self.cotizacion.inspeccion_origen_id, self.inspeccion.pk)

    def test_no_permite_dos_inspecciones_en_la_misma_recepcion(self):
        otra = self.InspeccionVehiculo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_inspeccion='INS-GRA-0002',
        )
        self._post(self._url(self.inspeccion), 'recepcion', str(self.recepcion.uuid))
        respuesta = self._post(self._url(otra), 'recepcion', str(self.recepcion.uuid))
        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        otra.refresh_from_db()
        self.assertIsNone(otra.recepcion_id)

    def test_no_vincula_cotizacion_ya_convertida(self):
        # La asociación operativa no autoriza editar la cotización desde la inspección.
        self.cotizacion.orden_trabajo_origen = self.orden
        self.cotizacion.save(update_fields=['orden_trabajo_origen', 'updated_at'])
        respuesta = self._post(self._url(), 'cotizacion', str(self.cotizacion.uuid))
        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.cotizacion.refresh_from_db()
        self.assertIsNone(self.cotizacion.inspeccion_origen_id)

    def test_no_edita_cita_desde_inspeccion(self):
        cita = self.Cita.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            fecha_cita=timezone.localdate(),
            hora_cita='09:30',
        )
        self._post(self._url(), 'recepcion', str(self.recepcion.uuid))
        respuesta = self._post(self._url(), 'cita', str(cita.uuid))
        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        cita.refresh_from_db()
        self.assertIsNone(cita.recepcion_generada_id)

    def test_no_vincula_cita_sin_recepcion_en_la_inspeccion(self):
        cita = self.Cita.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            fecha_cita=timezone.localdate(),
            hora_cita='09:30',
        )
        respuesta = self._post(self._url(), 'cita', str(cita.uuid))
        self.assertEqual(respuesta.status_code, 400, respuesta.content)

    def test_no_desvincula_orden_desde_inspeccion(self):
        self.inspeccion.orden_trabajo = self.orden
        self.inspeccion.save(update_fields=['orden_trabajo'])
        respuesta = self._post(self._url(), 'orden', str(self.orden.uuid), accion='desvincular')
        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.inspeccion.refresh_from_db()
        self.assertEqual(self.inspeccion.orden_trabajo_id, self.orden.pk)


class CotizacionRelacionesFlujoApiTests(_BaseRelacionesFlujoApiTests):
    def _url(self, cotizacion=None):
        return f'/api/cotizaciones/{(cotizacion or self.cotizacion).uuid}/relaciones/'

    def test_get_expone_recepcion_inspeccion_y_orden(self):
        respuesta = self.client.get(self._url())
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        datos = respuesta.json()
        self.assertEqual(set(datos['relaciones']), {'cita', 'recepcion', 'inspeccion', 'orden'})
        self.assertTrue(datos['puede_agregar']['recepcion'])

    def test_vincula_y_desvincula_recepcion(self):
        vincular = self._post(self._url(), 'recepcion', str(self.recepcion.uuid))
        self.assertEqual(vincular.status_code, 200, vincular.content)
        self.cotizacion.refresh_from_db()
        self.assertEqual(self.cotizacion.recepcion_origen_id, self.recepcion.pk)

        desvincular = self._post(self._url(), 'recepcion', str(self.recepcion.uuid), accion='desvincular')
        self.assertEqual(desvincular.status_code, 200, desvincular.content)
        self.cotizacion.refresh_from_db()
        self.assertIsNone(self.cotizacion.recepcion_origen_id)

    def test_asocia_orden_sin_incorporar_cotizacion(self):
        respuesta = self._post(self._url(), 'orden', str(self.orden.uuid))
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        self.cotizacion.refresh_from_db()
        self.orden.refresh_from_db()
        self.assertEqual(self.cotizacion.orden_trabajo_origen_id, self.orden.pk)
        self.assertIsNone(self.orden.cotizacion_origen_id)

        relaciones = self.client.get(self._url()).json()['relaciones']['orden']
        self.assertEqual([item['id'] for item in relaciones], [str(self.orden.uuid)])

    def test_asocia_presupuesto_adicional_a_orden_con_cotizacion_inicial(self):
        self.orden.cotizacion_origen = self.cotizacion
        self.orden.save(update_fields=['cotizacion_origen'])
        otra = self._nueva_cotizacion('COT-GRA-0002')
        respuesta = self._post(self._url(otra), 'orden', str(self.orden.uuid))
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        otra.refresh_from_db()
        self.assertEqual(otra.orden_trabajo_origen_id, self.orden.pk)

    def test_muestra_recepcion_heredada_de_la_inspeccion_sin_borrarla(self):
        self.cotizacion.inspeccion_origen = self.inspeccion
        self.inspeccion.recepcion = self.recepcion
        self.inspeccion.save(update_fields=['recepcion'])
        self.cotizacion.save(update_fields=['inspeccion_origen'])

        datos = self.client.get(self._url()).json()
        recepcion = datos['relaciones']['recepcion'][0]
        self.assertEqual(recepcion['id'], str(self.recepcion.uuid))
        self.assertFalse(recepcion['canDelete'])
        self.assertFalse(datos['puede_agregar']['recepcion'])

    def test_no_quita_el_vinculo_de_una_cotizacion_convertida(self):
        self.orden.cotizacion_origen = self.cotizacion
        self.orden.save(update_fields=['cotizacion_origen'])
        respuesta = self._post(self._url(), 'orden', str(self.orden.uuid), accion='desvincular')
        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.orden.refresh_from_db()
        self.assertEqual(self.orden.cotizacion_origen_id, self.cotizacion.pk)

    def _nueva_cotizacion(self, numero):
        from apps.cotizaciones.models import Cotizacion

        return Cotizacion.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_cotizacion=numero,
        )


class OrdenRelacionesFlujoApiTests(_BaseRelacionesFlujoApiTests):
    def _url(self, orden=None):
        return f'/api/ordenes/ordenes-trabajo/{(orden or self.orden).uuid}/relaciones/'

    def test_get_deriva_recepciones_e_inspeccion(self):
        self.recepcion.orden_trabajo = self.orden
        self.recepcion.save(update_fields=['orden_trabajo'])
        self.inspeccion.recepcion = self.recepcion
        self.inspeccion.orden_trabajo = self.orden
        self.inspeccion.save(update_fields=['recepcion', 'orden_trabajo'])

        datos = self.client.get(self._url()).json()
        self.assertEqual(
            [item['id'] for item in datos['relaciones']['recepcion']], [str(self.recepcion.uuid)]
        )
        self.assertEqual(
            [item['id'] for item in datos['relaciones']['inspeccion']], [str(self.inspeccion.uuid)]
        )
        self.assertFalse(any(datos['puede_agregar'].values()))

    def test_no_rep_recepcion_que_ya_tiene_otra_orden(self):
        otra = self.OrdenTrabajo.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_orden='OT-GRA-0002',
        )
        self.recepcion.orden_trabajo = otra
        self.recepcion.save(update_fields=['orden_trabajo'])
        respuesta = self._post(self._url(), 'recepcion', str(self.recepcion.uuid))
        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.recepcion.refresh_from_db()
        self.assertEqual(self.recepcion.orden_trabajo_id, otra.pk)

    def test_no_edita_cita_desde_orden(self):
        self.recepcion.orden_trabajo = self.orden
        self.recepcion.save(update_fields=['orden_trabajo'])
        cita = self.Cita.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            fecha_cita=timezone.localdate(),
            hora_cita='11:00',
        )
        respuesta = self._post(self._url(), 'cita', str(cita.uuid))
        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        cita.refresh_from_db()
        self.assertIsNone(cita.recepcion_generada_id)

    def test_no_vincula_cita_si_la_orden_no_tiene_recepcion(self):
        cita = self.Cita.objects.create(
            empresa=self.empresa,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            fecha_cita=timezone.localdate(),
            hora_cita='11:00',
        )
        respuesta = self._post(self._url(), 'cita', str(cita.uuid))
        self.assertEqual(respuesta.status_code, 400, respuesta.content)

    def test_rechaza_relacion_con_otra_empresa(self):
        from apps.clientes.models import Cliente
        from apps.empresas.models import Empresa

        otra_empresa = Empresa.objects.create(nombre_comercial='Taller Otro', ruc='999999999999')
        otro_cliente = Cliente.objects.create(empresa=otra_empresa, nombre='Cliente Otro')
        cotizacion = self._cotizacion_externa(otra_empresa, otro_cliente)
        respuesta = self._post(self._url(), 'cotizacion', str(cotizacion.uuid))
        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.orden.refresh_from_db()
        self.assertIsNone(self.orden.cotizacion_origen_id)

    def test_rechaza_tipo_y_accion_invalidos(self):
        respuesta = self._post(self._url(), 'vehiculo', str(self.vehiculo.uuid))
        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        respuesta = self.client.post(
            self._url(),
            {'tipo': 'cotizacion', 'id': str(self.cotizacion.uuid), 'accion': 'borrar'},
            format='json',
        )
        self.assertEqual(respuesta.status_code, 400, respuesta.content)

    def _cotizacion_externa(self, empresa, cliente):
        from apps.cotizaciones.models import Cotizacion

        return Cotizacion.objects.create(
            empresa=empresa, cliente=cliente, numero_cotizacion='COT-EXT-0001'
        )


class VinculoCotizacionFlujoTests(TestCase):
    """La cotización creada sin origen (WhatsApp, teléfono) se ata sola a la
    recepción que nace después. La inspección ya no hereda ítems de la
    cotización: se capturan directamente en la inspección."""

    @classmethod
    def setUpTestData(cls):
        from apps.clientes.models import Cliente
        from apps.cotizaciones.models import Cotizacion
        from apps.empresas.models import Empresa, Taller
        from apps.vehiculos.models import Vehiculo

        from django.contrib.auth import get_user_model

        User = get_user_model()
        cls.user = User.objects.create_superuser(username='flujo_admin', password='x')
        cls.empresa = Empresa.objects.create(
            nombre_comercial='Taller Flujo', razon_social='FLUJO SA', ruc='555555555555'
        )
        cls.taller = Taller.objects.create(
            empresa=cls.empresa, nombre='Central', direccion='Av. Siempre Viva 742'
        )
        cls.cliente = Cliente.objects.create(
            empresa=cls.empresa, nombre='Ana Lotus', identificacion='0912345678'
        )
        cls.vehiculo = Vehiculo.objects.create(placa='FLU-001', marca='Toyota', modelo='Corolla')
        cls.vehiculo.empresas.add(cls.empresa)
        cls.Cotizacion = Cotizacion

    def setUp(self):
        from rest_framework.test import APIClient

        self.api = APIClient()
        self.api.force_authenticate(user=self.user)
        fijar_empresa(self.api, str(self.empresa.uuid))
    def _recepcion(self, estado='ACEPTADA', numero='REC-0001'):
        from .models import RecepcionVehiculo

        return RecepcionVehiculo.objects.create(
            empresa=self.empresa,
            sucursal=self.taller,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_recepcion=numero,
            estado=estado,
            motivo_ingreso='Ruido en el motor',
        )

    def _cotizacion(self, estado='PENDIENTE', numero='COT-0001'):
        return self.Cotizacion.objects.create(
            empresa=self.empresa,
            sucursal=self.taller,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_cotizacion=numero,
            estado=estado,
        )

    # --- Recepción ---

    def test_recepcion_se_vincula_con_la_cotizacion_vigente(self):
        from apps.cotizaciones.services import conectar_recepcion

        cotizacion = self._cotizacion()
        recepcion = self._recepcion()

        vinculada = conectar_recepcion(recepcion)

        self.assertEqual(vinculada.pk, cotizacion.pk)
        cotizacion.refresh_from_db()
        self.assertEqual(cotizacion.recepcion_origen_id, recepcion.id)

    def test_recepcion_no_pisa_un_vinculo_existente(self):
        from apps.cotizaciones.services import conectar_recepcion

        primera = self._recepcion(numero='REC-0001')
        cotizacion = self._cotizacion()
        conectar_recepcion(primera)

        segunda = self._recepcion(numero='REC-0002')
        conectar_recepcion(segunda)

        cotizacion.refresh_from_db()
        self.assertEqual(cotizacion.recepcion_origen_id, primera.id)

    def test_sin_cotizacion_vigente_no_hace_nada(self):
        from apps.cotizaciones.services import conectar_recepcion

        recepcion = self._recepcion()
        self.assertIsNone(conectar_recepcion(recepcion))
        self.assertEqual(self.Cotizacion.objects.count(), 0)

    def test_cotizacion_vencida_no_se_vincula(self):
        from apps.cotizaciones.services import conectar_recepcion

        cotizacion = self._cotizacion(estado='VENCIDA', numero='COT-0009')
        recepcion = self._recepcion()

        self.assertIsNone(conectar_recepcion(recepcion))
        cotizacion.refresh_from_db()
        self.assertIsNone(cotizacion.recepcion_origen_id)


class InspeccionReaperturaApiTests(TestCase):
    """Una inspección FINALIZADA solo se reabre si el trabajo no avanzó: nada de
    orden de trabajo ni de cotización enviada/aceptada."""

    @classmethod
    def setUpTestData(cls):
        from apps.clientes.models import Cliente
        from apps.cotizaciones.models import Cotizacion
        from apps.empresas.models import Empresa, Taller
        from apps.vehiculos.models import Vehiculo

        from django.contrib.auth import get_user_model

        User = get_user_model()
        cls.user = User.objects.create_superuser(username='reapertura_admin', password='x')
        cls.empresa = Empresa.objects.create(
            nombre_comercial='Taller Reapertura', razon_social='REAP SA', ruc='999999999999'
        )
        cls.taller = Taller.objects.create(
            empresa=cls.empresa, nombre='Central', direccion='Av. Siempre Viva 742'
        )
        cls.cliente = Cliente.objects.create(
            empresa=cls.empresa, nombre='Ana Lotus', identificacion='0912345678'
        )
        cls.vehiculo = Vehiculo.objects.create(placa='REA-001', marca='Toyota', modelo='Corolla')
        cls.vehiculo.empresas.add(cls.empresa)
        cls.Cotizacion = Cotizacion

    def setUp(self):
        from rest_framework.test import APIClient

        from .models import InspeccionVehiculo

        self.api = APIClient()
        self.api.force_authenticate(user=self.user)
        fijar_empresa(self.api, str(self.empresa.uuid))
        self.inspeccion = InspeccionVehiculo.objects.create(
            empresa=self.empresa,
            sucursal=self.taller,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_inspeccion='INS-REA-0001',
            estado='FINALIZADA',
        )

    def _url(self):
        return f'/api/ordenes/inspecciones/{self.inspeccion.uuid}/'

    def _cotizacion(self, estado):
        return self.Cotizacion.objects.create(
            empresa=self.empresa,
            sucursal=self.taller,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            inspeccion_origen=self.inspeccion,
            numero_cotizacion=f'COT-REA-{estado}',
            estado=estado,
        )

    def test_reabre_sin_bloqueos(self):
        respuesta = self.api.patch(self._url(), {'estado': 'EN_PROCESO'}, format='json')

        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        self.inspeccion.refresh_from_db()
        self.assertEqual(self.inspeccion.estado, 'EN_PROCESO')
        self.assertIsNone(self.inspeccion.fecha_finalizacion)

    def test_no_reabre_con_orden_de_trabajo(self):
        from .models import OrdenTrabajo

        orden = OrdenTrabajo.objects.create(
            empresa=self.empresa,
            sucursal=self.taller,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            numero_orden='OT-REA-0001',
        )
        self.inspeccion.orden_trabajo = orden
        self.inspeccion.save(update_fields=['orden_trabajo', 'updated_at'])

        respuesta = self.api.patch(self._url(), {'estado': 'EN_PROCESO'}, format='json')

        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.assertIn('orden de trabajo', str(respuesta.data).lower())
        self.inspeccion.refresh_from_db()
        self.assertEqual(self.inspeccion.estado, 'FINALIZADA')

    def test_no_reabre_con_cotizacion_enviada(self):
        self._cotizacion('ENVIADA')

        respuesta = self.api.patch(self._url(), {'estado': 'EN_PROCESO'}, format='json')

        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.inspeccion.refresh_from_db()
        self.assertEqual(self.inspeccion.estado, 'FINALIZADA')

    def test_no_reabre_con_cotizacion_aceptada(self):
        self._cotizacion('ACEPTADA')

        respuesta = self.api.patch(self._url(), {'estado': 'EN_PROCESO'}, format='json')

        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.inspeccion.refresh_from_db()
        self.assertEqual(self.inspeccion.estado, 'FINALIZADA')

    def test_reabre_con_cotizacion_pendiente(self):
        self._cotizacion('PENDIENTE')

        respuesta = self.api.patch(self._url(), {'estado': 'EN_PROCESO'}, format='json')

        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        self.inspeccion.refresh_from_db()
        self.assertEqual(self.inspeccion.estado, 'EN_PROCESO')

    def test_detalle_marca_puede_reabrir(self):
        respuesta = self.api.get(self._url())

        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        self.assertTrue(respuesta.data['puede_reabrir'])
        self.assertFalse(respuesta.data['tiene_cotizacion_bloqueante'])

    def test_detalle_bloquea_reapertura_con_cotizacion_enviada(self):
        self._cotizacion('ENVIADA')

        respuesta = self.api.get(self._url())

        self.assertFalse(respuesta.data['puede_reabrir'])
        self.assertTrue(respuesta.data['tiene_cotizacion_bloqueante'])


class RecepcionCreadaConTallerActivoTests(TestCase):
    """Al crear una recepción, la sucursal debe ser el taller activo de la sesión
    (``X-Taller-ID``), no el primer taller de la empresa. Si se guarda con otro
    taller, el filtro estricto por taller la excluye y el detalle recién creado
    devuelve 404 ("No RecepcionVehiculo matches the given query")."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from rest_framework.test import APIClient

        from apps.clientes.models import Cliente
        from apps.empresas.models import Empresa, Taller
        from apps.vehiculos.models import Vehiculo
        from .models import RecepcionVehiculo, TipoTrabajo

        self.empresa = Empresa.objects.create(
            nombre_comercial='Taller Activo', ruc='444444444444'
        )
        # El primer taller (por id) NO es el activo: reproduce el bug original.
        self.taller_primero = Taller.objects.create(
            empresa=self.empresa, nombre='Taller Uno',
            codigo_sucursal='001', prefijo_recepcion='REC-',
        )
        self.taller_activo = Taller.objects.create(
            empresa=self.empresa, nombre='Taller Dos',
            codigo_sucursal='002', prefijo_recepcion='REC-',
        )
        self.cliente = Cliente.objects.create(
            empresa=self.empresa, nombre='Cliente Activo', identificacion='1712345678'
        )
        self.vehiculo = Vehiculo.objects.create(
            placa='ACT-1234', marca='Chevrolet', modelo='Aveo'
        )
        self.vehiculo.empresas.add(self.empresa)

        usuario = get_user_model().objects.create_superuser(
            username='taller-activo', email='taller-activo@test.local', password='x'
        )
        self.client = APIClient()
        self.client.force_authenticate(usuario)
        fijar_empresa(self.client, str(self.empresa.uuid))
        self.client.defaults['HTTP_X_TALLER_ID'] = str(self.taller_activo.uuid)
        self.tipo = TipoTrabajo.MANTENIMIENTO

    def _payload(self):
        return {
            'cliente': str(self.cliente.uuid),
            'vehiculo': str(self.vehiculo.uuid),
            'tipo_recepcion': self.tipo,
            'estado': 'NO_ACEPTADA',
            'motivo_no_recepcion': 'El cliente no desea firmar',
        }

    def test_crea_con_el_taller_activo(self):
        from .models import RecepcionVehiculo

        respuesta = self.client.post('/api/recepciones/', self._payload(), format='json')

        self.assertEqual(respuesta.status_code, 201, respuesta.content)
        creada = RecepcionVehiculo.objects.get(uuid=respuesta.data['id'])
        self.assertEqual(creada.sucursal_id, self.taller_activo.pk)
        self.assertNotEqual(creada.sucursal_id, self.taller_primero.pk)

    def test_crear_por_multipart_queda_activa(self):
        """El formulario envía FormData sin `is_active`: no debe leerse como False."""
        from .models import RecepcionVehiculo

        respuesta = self.client.post('/api/recepciones/', self._payload(), format='multipart')

        self.assertEqual(respuesta.status_code, 201, respuesta.content)
        self.assertTrue(RecepcionVehiculo.objects.get(uuid=respuesta.data['id']).is_active)

    def test_el_detalle_del_registro_creado_es_visible_en_el_taller_activo(self):
        creada = self.client.post('/api/recepciones/', self._payload(), format='json').data

        detalle = self.client.get(f"/api/recepciones/{creada['id']}/")

        self.assertEqual(detalle.status_code, 200, detalle.content)


class RecepcionFirmarApiTests(TestCase):
    """Re-firma de una recepción NO_ACEPTADA: el cliente cambia de opinión y la
    firma la deja directamente como ACEPTADA (nunca pasa por PENDIENTE)."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from django.utils import timezone
        from rest_framework.test import APIClient

        from apps.clientes.models import Cliente
        from apps.empresas.models import Empresa, Taller
        from apps.vehiculos.models import Vehiculo
        from .models import RecepcionVehiculo, TipoTrabajo

        self.empresa = Empresa.objects.create(
            nombre_comercial='Taller Firma', ruc='555555555555'
        )
        self.taller = Taller.objects.create(
            empresa=self.empresa, nombre='Taller Firma',
            codigo_sucursal='001', prefijo_recepcion='REC-',
        )
        self.cliente = Cliente.objects.create(
            empresa=self.empresa, nombre='Cliente Firma', identificacion='1799999999'
        )
        self.vehiculo = Vehiculo.objects.create(
            placa='FIR-1234', marca='Chevrolet', modelo='Spark'
        )
        self.vehiculo.empresas.add(self.empresa)

        self.recepcion = RecepcionVehiculo.objects.create(
            empresa=self.empresa,
            sucursal=self.taller,
            cliente=self.cliente,
            vehiculo=self.vehiculo,
            tipo_recepcion=TipoTrabajo.MANTENIMIENTO,
            estado='NO_ACEPTADA',
            motivo_no_recepcion='El cliente no firmó',
        )

        usuario = get_user_model().objects.create_superuser(
            username='firma-test', email='firma-test@test.local', password='x'
        )
        self.client = APIClient()
        self.client.force_authenticate(usuario)
        fijar_empresa(self.client, str(self.empresa.uuid))
        self.client.defaults['HTTP_X_TALLER_ID'] = str(self.taller.uuid)

    def _url(self, pk=None):
        return f'/api/recepciones/{pk or self.recepcion.uuid}/firmar/'

    def test_firma_no_aceptada_pasa_directo_a_aceptada(self):
        respuesta = self.client.post(
            self._url(), {'firma_cliente': firma_data_url()}, format='json'
        )

        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        self.recepcion.refresh_from_db()
        self.assertEqual(self.recepcion.estado, 'ACEPTADA')
        self.assertTrue(self.recepcion.aceptacion_condiciones)
        self.assertIsNotNone(self.recepcion.fecha_firma_cliente)
        # La firma se guarda como archivo PNG y la API devuelve su URL.
        self.assertTrue(self.recepcion.firma_cliente.name.endswith('.png'))
        self.assertTrue(respuesta.data['firma_cliente'].startswith('http'))
        self.assertFalse(self.recepcion.motivo_no_recepcion)
        self.assertEqual(respuesta.data['estado'], 'ACEPTADA')

    def test_firmar_acepta_un_archivo_png(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        archivo = SimpleUploadedFile('firma.png', imagen_de_prueba(400, 100), content_type='image/png')
        respuesta = self.client.post(self._url(), {'firma_cliente': archivo}, format='multipart')

        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        self.recepcion.refresh_from_db()
        self.assertTrue(self.recepcion.firma_cliente.name.endswith('.png'))

    def test_rechaza_una_firma_que_no_es_imagen(self):
        respuesta = self.client.post(
            self._url(), {'firma_cliente': 'data:image/png;base64,AAAA'}, format='json'
        )

        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.assertIn('firma_cliente', respuesta.data)

    def test_no_permite_firmar_sin_firma(self):
        respuesta = self.client.post(self._url(), {}, format='json')

        self.assertEqual(respuesta.status_code, 400, respuesta.content)
        self.recepcion.refresh_from_db()
        self.assertEqual(self.recepcion.estado, 'NO_ACEPTADA')

    def test_no_permite_firmar_una_recepcion_ya_aceptada(self):
        self.recepcion.estado = 'ACEPTADA'
        self.recepcion.aceptacion_condiciones = True
        self.recepcion.save(update_fields=['estado', 'aceptacion_condiciones'])

        respuesta = self.client.post(
            self._url(), {'firma_cliente': firma_data_url()}, format='json'
        )

        self.assertEqual(respuesta.status_code, 400, respuesta.content)


    def test_firmar_y_crear_inspeccion_desde_la_recepcion(self):
        from .models import InspeccionVehiculo

        firma = self.client.post(
            self._url(), {'firma_cliente': firma_data_url()}, format='json'
        )
        self.assertEqual(firma.status_code, 200, firma.content)
        type(self.recepcion).objects.filter(pk=self.recepcion.pk).update(motivo_ingreso='Mantenimiento de 10.000 km')

        respuesta = self.client.post(f'/api/recepciones/{self.recepcion.uuid}/crear-inspeccion/')

        self.assertEqual(respuesta.status_code, 201, respuesta.content)
        inspeccion = InspeccionVehiculo.objects.get(uuid=respuesta.json()['id'])
        self.assertEqual(inspeccion.recepcion_id, self.recepcion.pk)
        self.assertEqual(respuesta.json()['recepcion']['id'], str(self.recepcion.uuid))

    def test_testigos_ev_y_adas_pasan_de_la_recepcion_a_la_inspeccion(self):
        from .models import InspeccionVehiculo

        firma = self.client.post(
            self._url(), {'firma_cliente': firma_data_url()}, format='json'
        )
        self.assertEqual(firma.status_code, 200, firma.content)
        type(self.recepcion).objects.filter(pk=self.recepcion.pk).update(
            motivo_ingreso='Revisión de testigos', testigo_ev_listo=True, testigo_dpf=True, testigo_punto_ciego=True,
        )
        recepcion = self.client.get(f'/api/recepciones/{self.recepcion.uuid}/')
        self.assertEqual(recepcion.status_code, 200, recepcion.content)
        self.assertTrue(recepcion.json()['testigo_dpf'])
        self.assertFalse(recepcion.json()['testigo_precolision'])

        respuesta = self.client.post(f'/api/recepciones/{self.recepcion.uuid}/crear-inspeccion/')

        self.assertEqual(respuesta.status_code, 201, respuesta.content)
        datos = respuesta.json()
        self.assertTrue(datos['testigo_ev_listo'])
        self.assertTrue(datos['testigo_dpf'])
        self.assertTrue(datos['testigo_punto_ciego'])
        self.assertFalse(datos['testigo_salida_carril'])

        edicion = self.client.patch(
            f'/api/ordenes/inspecciones/{datos["id"]}/', {'testigo_salida_carril': True}, format='json'
        )
        self.assertEqual(edicion.status_code, 200, edicion.content)
        inspeccion = InspeccionVehiculo.objects.get(uuid=datos['id'])
        self.assertTrue(inspeccion.testigo_salida_carril)
        self.assertTrue(inspeccion.testigo_punto_ciego)

class BackfillSucursalDocumentosTests(TestCase):
    """La migración de backfill asigna taller a los documentos huérfanos
    (creados cuando el API no persistía la sucursal)."""

    def _run_migration(self, modulo, nombre):
        import importlib

        from django.apps import apps as django_apps

        migracion = importlib.import_module(modulo)
        getattr(migracion, nombre)(django_apps, None)

    def setUp(self):
        from apps.clientes.models import Cliente
        from apps.empresas.models import Empresa, Taller
        from apps.vehiculos.models import Vehiculo
        from .models import InspeccionVehiculo, OrdenTrabajo, RecepcionVehiculo

        self.empresa = Empresa.objects.create(
            nombre_comercial='Backfill', ruc='333333333333'
        )
        self.taller_a = Taller.objects.create(
            empresa=self.empresa, nombre='Taller A', codigo_sucursal='001',
        )
        self.taller_b = Taller.objects.create(
            empresa=self.empresa, nombre='Taller B', codigo_sucursal='002',
        )
        self.cliente = Cliente.objects.create(
            empresa=self.empresa, nombre='Cliente Huérfano'
        )
        self.vehiculo = Vehiculo.objects.create(placa='ORF-0001', marca='Kia', modelo='Rio')

        self.recepcion = RecepcionVehiculo.objects.create(
            empresa=self.empresa, sucursal=None, cliente=self.cliente,
            vehiculo=self.vehiculo, numero_recepcion='REC-00001',
        )
        self.inspeccion_con_recepcion = InspeccionVehiculo.objects.create(
            empresa=self.empresa, sucursal=None, cliente=self.cliente,
            vehiculo=self.vehiculo, recepcion=self.recepcion,
            numero_inspeccion='INS-00001',
        )
        self.inspeccion_suelta = InspeccionVehiculo.objects.create(
            empresa=self.empresa, sucursal=None, cliente=self.cliente,
            vehiculo=self.vehiculo, recepcion=None, numero_inspeccion='INS-00002',
        )
        self.orden = OrdenTrabajo.objects.create(
            empresa=self.empresa, sucursal=None, cliente=self.cliente,
            vehiculo=self.vehiculo, numero_orden='OT-00001',
        )

    def test_backfill_ordenes_asigna_taller_por_defecto_y_hereda_recepcion(self):
        from .models import RecepcionVehiculo

        self._run_migration(
            'apps.ordenes.migrations.0046_backfill_sucursal_documentos',
            'asignar_sucursal',
        )

        for documento in (self.recepcion, self.inspeccion_suelta, self.orden):
            documento.refresh_from_db()
            self.assertEqual(documento.sucursal_id, self.taller_a.pk)
        self.inspeccion_con_recepcion.refresh_from_db()
        self.assertEqual(self.inspeccion_con_recepcion.sucursal_id, self.taller_a.pk)

    def test_backfill_cotizacion_hereda_del_origen(self):
        from apps.cotizaciones.models import Cotizacion

        # El origen ya tiene taller B: la cotización debe heredarlo, no el default.
        self.recepcion.sucursal = self.taller_b
        self.recepcion.save(update_fields=['sucursal'])
        cotizacion = Cotizacion.objects.create(
            empresa=self.empresa, sucursal=None, cliente=self.cliente,
            vehiculo=self.vehiculo, recepcion_origen=self.recepcion,
            numero_cotizacion='COT-00001',
        )

        self._run_migration(
            'apps.cotizaciones.migrations.0019_backfill_sucursal_desde_origen',
            'asignar_sucursal',
        )

        cotizacion.refresh_from_db()
        self.assertEqual(cotizacion.sucursal_id, self.taller_b.pk)


class OrdenSucursalConservadaTests(TestCase):
    """La orden toma el taller de ``X-Taller-ID`` al crearse y una edición con
    ``sucursal: null`` no lo borra (si no, el filtro estricto la ocultaría)."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from rest_framework.test import APIClient

        from apps.clientes.models import Cliente
        from apps.empresas.models import Empresa, Taller
        from apps.vehiculos.models import Vehiculo

        self.empresa = Empresa.objects.create(nombre_comercial='Taller OT Sucursal', ruc='555555555555')
        Taller.objects.create(empresa=self.empresa, nombre='Taller Uno', codigo_sucursal='001', prefijo_ot='OT-')
        self.taller_activo = Taller.objects.create(
            empresa=self.empresa, nombre='Taller Dos', codigo_sucursal='002', prefijo_ot='OT-'
        )
        self.cliente = Cliente.objects.create(
            empresa=self.empresa, nombre='Cliente OT', identificacion='1712345679'
        )
        self.vehiculo = Vehiculo.objects.create(placa='OTS-1234', marca='Kia', modelo='Rio')
        self.vehiculo.empresas.add(self.empresa)

        usuario = get_user_model().objects.create_superuser(
            username='ot-sucursal', email='ot-sucursal@test.local', password='x'
        )
        self.responsable = str(usuario.profile.uuid)
        self.client = APIClient()
        self.client.force_authenticate(usuario)
        fijar_empresa(self.client, str(self.empresa.uuid))
        self.client.defaults['HTTP_X_TALLER_ID'] = str(self.taller_activo.uuid)

    def _payload(self, **extra):
        return {
            'sucursal': None,
            'cliente': str(self.cliente.uuid),
            'vehiculo': str(self.vehiculo.uuid),
            'estado': 'PENDIENTE',
            'prioridad': 'MEDIA',
            'tipo_trabajo': 'MANTENIMIENTO',
            'kilometraje': self.vehiculo.kilometraje_actual,
            'fecha_ingreso': '2026-10-01T08:30:00Z',
            'mecanico_principal': self.responsable,
            **extra,
        }

    def test_crear_sin_sucursal_usa_el_taller_activo(self):
        from .models import OrdenTrabajo

        resp = self.client.post('/api/ordenes/ordenes-trabajo/', self._payload(), format='json')

        self.assertEqual(resp.status_code, 201, resp.data)
        self.assertEqual(OrdenTrabajo.objects.get(uuid=resp.data['id']).sucursal_id, self.taller_activo.pk)

    def test_fecha_ingreso_editable(self):
        """La fecha de ingreso la registra el usuario al crear y al editar."""
        from datetime import datetime, timezone as dt_timezone

        from .models import OrdenTrabajo

        al_crear = datetime(2026, 10, 1, 8, 30, tzinfo=dt_timezone.utc)
        resp = self.client.post(
            '/api/ordenes/ordenes-trabajo/',
            self._payload(fecha_ingreso=al_crear.isoformat()),
            format='json',
        )
        self.assertEqual(resp.status_code, 201, resp.data)
        orden = OrdenTrabajo.objects.get(uuid=resp.data['id'])
        self.assertEqual(orden.fecha_ingreso, al_crear)

        al_editar = datetime(2026, 10, 2, 14, 15, tzinfo=dt_timezone.utc)
        resp = self.client.patch(
            f'/api/ordenes/ordenes-trabajo/{orden.uuid}/',
            {'fecha_ingreso': al_editar.isoformat()},
            format='json',
        )
        self.assertEqual(resp.status_code, 200, resp.data)
        orden.refresh_from_db()
        self.assertEqual(orden.fecha_ingreso, al_editar)

    def test_fecha_ingreso_obligatoria(self):
        """Sin fecha de ingreso no se crea la orden, y al editar no puede quedar vacía."""
        payload = self._payload()
        payload.pop('fecha_ingreso')
        resp = self.client.post('/api/ordenes/ordenes-trabajo/', payload, format='json')
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(str(resp.data['fecha_ingreso'][0]), 'La fecha de ingreso es obligatoria.')

        creada = self.client.post('/api/ordenes/ordenes-trabajo/', self._payload(), format='json').data
        for vacio in (None, ''):
            resp = self.client.patch(
                f"/api/ordenes/ordenes-trabajo/{creada['id']}/",
                {'fecha_ingreso': vacio},
                format='json',
            )
            self.assertEqual(resp.status_code, 400, vacio)
            self.assertIn('fecha_ingreso', resp.data)

        # Un PATCH que no envía el campo (p. ej. cambio de estado) conserva la fecha.
        resp = self.client.patch(
            f"/api/ordenes/ordenes-trabajo/{creada['id']}/",
            {'prioridad': 'ALTA'},
            format='json',
        )
        self.assertEqual(resp.status_code, 200, resp.data)

    def test_responsable_obligatorio(self):
        """Sin responsable no se crea la orden, y al editar no puede quedar vacío."""
        payload = self._payload()
        payload.pop('mecanico_principal')
        resp = self.client.post('/api/ordenes/ordenes-trabajo/', payload, format='json')
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(str(resp.data['mecanico_principal'][0]), 'El responsable de la orden es obligatorio.')

        creada = self.client.post('/api/ordenes/ordenes-trabajo/', self._payload(), format='json').data
        for vacio in (None, ''):
            resp = self.client.patch(
                f"/api/ordenes/ordenes-trabajo/{creada['id']}/",
                {'mecanico_principal': vacio},
                format='json',
            )
            self.assertEqual(resp.status_code, 400, vacio)
            self.assertIn('mecanico_principal', resp.data)

        # Un PATCH que no envía el campo (p. ej. cambio de estado) conserva el responsable.
        resp = self.client.patch(
            f"/api/ordenes/ordenes-trabajo/{creada['id']}/",
            {'prioridad': 'ALTA'},
            format='json',
        )
        self.assertEqual(resp.status_code, 200, resp.data)

    def test_editar_con_sucursal_null_conserva_el_taller(self):
        from .models import OrdenTrabajo

        creada = self.client.post('/api/ordenes/ordenes-trabajo/', self._payload(), format='json').data

        resp = self.client.put(
            f"/api/ordenes/ordenes-trabajo/{creada['id']}/",
            self._payload(observaciones='Editada'),
            format='json',
        )

        self.assertEqual(resp.status_code, 200, resp.data)
        orden = OrdenTrabajo.objects.get(uuid=creada['id'])
        self.assertEqual(orden.sucursal_id, self.taller_activo.pk)
        self.assertEqual(orden.observaciones, 'Editada')


class SalidaRecepcionAlEntregarOrdenTests(TestCase):
    """Al pasar la orden a ENTREGADO, la recepción del flujo toma como fecha de
    salida la fecha de entrega de la orden (signal `post_save`)."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from rest_framework.test import APIClient

        from apps.clientes.models import Cliente
        from apps.empresas.models import Empresa, Taller
        from apps.vehiculos.models import Vehiculo

        from .models import OrdenTrabajo, RecepcionVehiculo

        self.empresa = Empresa.objects.create(nombre_comercial='Taller Entrega', ruc='666666666666')
        self.taller = Taller.objects.create(
            empresa=self.empresa, nombre='Taller Entrega', codigo_sucursal='001', prefijo_ot='OT-'
        )
        cliente = Cliente.objects.create(empresa=self.empresa, nombre='Cliente Entrega', identificacion='1712345680')
        vehiculo = Vehiculo.objects.create(placa='ENT-1234', marca='Kia', modelo='Rio')
        vehiculo.empresas.add(self.empresa)

        self.orden = OrdenTrabajo.objects.create(
            empresa=self.empresa, sucursal=self.taller, cliente=cliente, vehiculo=vehiculo,
            numero_orden='OT-ENT-1', estado='EN_PROCESO',
        )
        self.recepcion = RecepcionVehiculo.objects.create(
            empresa=self.empresa, sucursal=self.taller, cliente=cliente, vehiculo=vehiculo,
            numero_recepcion='REC-ENT-1', estado='ACEPTADA', orden_trabajo=self.orden,
        )
        self.inspeccion_recepcion = RecepcionVehiculo.objects.create(
            empresa=self.empresa, sucursal=self.taller, cliente=cliente, vehiculo=vehiculo,
            numero_recepcion='REC-ENT-2', estado='ACEPTADA',
        )
        self.otra_recepcion = RecepcionVehiculo.objects.create(
            empresa=self.empresa, sucursal=self.taller, cliente=cliente, vehiculo=vehiculo,
            numero_recepcion='REC-ENT-3', estado='ACEPTADA',
        )
        InspeccionVehiculo.objects.create(
            empresa=self.empresa, sucursal=self.taller, recepcion=self.inspeccion_recepcion,
            cliente=cliente, vehiculo=vehiculo, numero_inspeccion='INS-ENT-1',
            tipo_inspeccion='MANTENIMIENTO', orden_trabajo=self.orden,
        )

        usuario = get_user_model().objects.create_superuser(
            username='ot-entrega', email='ot-entrega@test.local', password='x'
        )
        self.client = APIClient()
        self.client.force_authenticate(usuario)
        fijar_empresa(self.client, str(self.empresa.uuid))
        self.client.defaults['HTTP_X_TALLER_ID'] = str(self.taller.uuid)

    def _entregar(self, **extra):
        return self.client.patch(
            f'/api/ordenes/ordenes-trabajo/{self.orden.uuid}/',
            {'estado': 'ENTREGADO', **extra},
            format='json',
        )

    def test_entregar_copia_la_fecha_de_entrega_a_las_recepciones_del_flujo(self):
        fecha = timezone.now().replace(microsecond=0) - timedelta(hours=2)
        resp = self._entregar(fecha_entrega=fecha.isoformat())
        self.assertEqual(resp.status_code, 200, resp.content)

        for recepcion in (self.recepcion, self.inspeccion_recepcion):
            recepcion.refresh_from_db()
            self.assertEqual(recepcion.fecha_salida, fecha)
        self.otra_recepcion.refresh_from_db()
        self.assertIsNone(self.otra_recepcion.fecha_salida)

    def test_entregar_sin_fecha_usa_ahora_en_orden_y_recepcion(self):
        resp = self._entregar(fecha_entrega=None)
        self.assertEqual(resp.status_code, 200, resp.content)

        self.orden.refresh_from_db()
        self.recepcion.refresh_from_db()
        self.assertIsNotNone(self.orden.fecha_entrega)
        self.assertEqual(self.recepcion.fecha_salida, self.orden.fecha_entrega)

    def test_corregir_fecha_de_entrega_actualiza_la_salida(self):
        self._entregar()
        nueva = timezone.now().replace(microsecond=0) - timedelta(days=1)
        resp = self.client.patch(
            f'/api/ordenes/ordenes-trabajo/{self.orden.uuid}/',
            {'fecha_entrega': nueva.isoformat()},
            format='json',
        )
        self.assertEqual(resp.status_code, 200, resp.content)
        self.recepcion.refresh_from_db()
        self.assertEqual(self.recepcion.fecha_salida, nueva)

    def test_orden_no_entregada_no_toca_la_recepcion(self):
        resp = self.client.patch(
            f'/api/ordenes/ordenes-trabajo/{self.orden.uuid}/',
            {'estado': 'COMPLETADO'},
            format='json',
        )
        self.assertEqual(resp.status_code, 200, resp.content)
        self.recepcion.refresh_from_db()
        self.assertIsNone(self.recepcion.fecha_salida)


class ExportarPdfInspeccionApiTests(_BaseRelacionesFlujoApiTests):
    """``GET /api/ordenes/inspecciones/<id>/exportar-pdf/``: informe para el cliente."""

    def _url(self, inspeccion=None):
        return f'/api/ordenes/inspecciones/{(inspeccion or self.inspeccion).uuid}/exportar-pdf/'

    def _assert_pdf(self, respuesta, numero='INS-GRA-0001'):
        self.assertEqual(respuesta.status_code, 200, getattr(respuesta, 'content', b''))
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertIn(f'inspeccion_{numero}.pdf', respuesta['Content-Disposition'])
        contenido = b''.join(respuesta.streaming_content)
        self.assertTrue(contenido.startswith(b'%PDF'))
        return contenido

    def test_descarga_informe_con_las_cuatro_fases(self):
        from django.core.files.base import ContentFile

        from .models import DetalleRepuestoInspeccion, DetalleServicioInspeccion, FotoInspeccion

        InspeccionVehiculo.objects.filter(pk=self.inspeccion.pk).update(
            motivo_ingreso='Vibración al frenar',
            codigos_dtc='p0301, P0171;P0301',
            diagnostico_tecnico='Discos con alabeo.\nPastillas al límite.',
            recomendaciones='Cambiar discos y pastillas.',
            testigo_check_engine=True,
            testigo_abs=True,
            otros_testigos_observaciones='ABS intermitente',
            estado='FINALIZADA',
        )
        foto = FotoInspeccion(inspeccion=self.inspeccion, descripcion='Disco delantero')
        foto.imagen.save('disco.jpg', ContentFile(imagen_de_prueba(400, 300, 'JPEG')))
        FotoInspeccion.objects.create(inspeccion=self.inspeccion, imagen='no/existe.jpg')
        DetalleServicioInspeccion.objects.create(
            inspeccion=self.inspeccion, descripcion='Rectificado de discos', horas_estimadas='1.5', prioridad='ALTA',
        )
        DetalleRepuestoInspeccion.objects.create(
            inspeccion=self.inspeccion, descripcion='Pastillas', cantidad=1, prioridad='BAJA', es_sugerido=True,
        )
        self.cotizacion.inspeccion_origen = self.inspeccion
        self.cotizacion.save()

        self._assert_pdf(self.client.get(self._url()))

    def test_informe_sin_datos_no_falla(self):
        self._assert_pdf(self.client.get(self._url()))

    def test_inspeccion_desde_recepcion_sin_cliente_propio(self):
        inspeccion = InspeccionVehiculo.objects.create(
            empresa=self.empresa, recepcion=self.recepcion, numero_inspeccion='INS-GRA-0002',
        )
        self._assert_pdf(self.client.get(self._url(inspeccion)), 'INS-GRA-0002')

    def test_texto_con_caracteres_de_markup_no_rompe_el_pdf(self):
        self.cliente.nombre = 'Pérez & Hijos <SA>'
        self.cliente.save()
        InspeccionVehiculo.objects.filter(pk=self.inspeccion.pk).update(
            diagnostico_tecnico='Presión < 20 psi & fuga', codigos_dtc='<b>P0300</b>',
        )
        self._assert_pdf(self.client.get(self._url()))

    def test_requiere_autenticacion(self):
        from rest_framework.test import APIClient

        self.assertEqual(APIClient().get(self._url()).status_code, 401)

    def test_no_filtra_inspecciones_de_otra_empresa(self):
        from apps.empresas.models import Empresa

        otra = Empresa.objects.create(nombre_comercial='Otra', ruc='777777777777')
        ajena = InspeccionVehiculo.objects.create(empresa=otra, numero_inspeccion='INS-AJENA')
        self.assertEqual(self.client.get(self._url(ajena)).status_code, 404)



class ExportarPdfOrdenApiTests(_BaseRelacionesFlujoApiTests):
    """``GET /api/ordenes/ordenes-trabajo/<id>/exportar-pdf/``: documento para el cliente."""

    def _url(self, orden=None):
        return f'/api/ordenes/ordenes-trabajo/{(orden or self.orden).uuid}/exportar-pdf/'

    def _assert_pdf(self, respuesta, numero='OT-GRA-0001'):
        self.assertEqual(respuesta.status_code, 200, getattr(respuesta, 'content', b''))
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertIn(f'orden_{numero}.pdf', respuesta['Content-Disposition'])
        contenido = b''.join(respuesta.streaming_content)
        self.assertTrue(contenido.startswith(b'%PDF'))
        return contenido

    def test_descarga_orden_con_trabajo_valores_y_fotos(self):
        from django.contrib.auth import get_user_model
        from django.core.files.base import ContentFile

        from .models import DetalleRepuestoOrdenTrabajo, DetalleServicioOrdenTrabajo, FotoOrdenTrabajo

        tecnico = get_user_model().objects.create_user(
            username='tecnico', first_name='Pedro', last_name='Jama', password='x',
        )
        DetalleServicioOrdenTrabajo.objects.create(
            orden_trabajo=self.orden, descripcion='Rectificado de discos', horas_aplicadas=Decimal('1.5'),
            precio_unitario=Decimal('20.00'), mecanico_asignado=tecnico, completado=True, servicio_codigo='SRV-01',
        )
        DetalleServicioOrdenTrabajo.objects.create(
            orden_trabajo=self.orden, descripcion='Alineación', precio_unitario=Decimal('15.00'),
            descuento=Decimal('5.00'),
        )
        DetalleRepuestoOrdenTrabajo.objects.create(
            orden_trabajo=self.orden, descripcion='Pastillas', precio_unitario=Decimal('48.50'),
            subtotal=Decimal('48.50'), iva_porcentaje=Decimal('0'),
        )
        self.orden.calcular_totales()
        foto = FotoOrdenTrabajo(orden_trabajo=self.orden, descripcion='Pastillas nuevas')
        foto.imagen.save('pastillas.jpg', ContentFile(imagen_de_prueba(400, 300, 'JPEG')))
        FotoOrdenTrabajo.objects.create(orden_trabajo=self.orden, imagen='no/existe.jpg')
        self.recepcion.orden_trabajo = self.orden
        self.recepcion.save()
        self.inspeccion.orden_trabajo = self.orden
        self.inspeccion.save()
        self.OrdenTrabajo.objects.filter(pk=self.orden.pk).update(
            estado='EN_ESPERA', motivo_espera='Repuestos', observaciones='Revisar suspensión.',
            cotizacion_origen=self.cotizacion, asesor=tecnico, mecanico_principal=tecnico,
        )

        self._assert_pdf(self.client.get(self._url()))

    def test_orden_sin_items_ni_fotos_no_falla(self):
        self._assert_pdf(self.client.get(self._url()))

    def test_orden_entregada(self):
        from django.utils import timezone

        self.OrdenTrabajo.objects.filter(pk=self.orden.pk).update(estado='ENTREGADO', fecha_entrega=timezone.now())
        self._assert_pdf(self.client.get(self._url()))

    def test_texto_con_caracteres_de_markup_no_rompe_el_pdf(self):
        from .models import DetalleServicioOrdenTrabajo

        self.cliente.nombre = 'Pérez & Hijos <SA>'
        self.cliente.save()
        DetalleServicioOrdenTrabajo.objects.create(
            orden_trabajo=self.orden, descripcion='Presión < 20 psi & <b>fuga</b>', precio_unitario=Decimal('10.00'),
        )
        self.OrdenTrabajo.objects.filter(pk=self.orden.pk).update(observaciones='<i>Nota</i> & más')
        self._assert_pdf(self.client.get(self._url()))

    def test_requiere_autenticacion(self):
        from rest_framework.test import APIClient

        self.assertEqual(APIClient().get(self._url()).status_code, 401)

    def test_no_filtra_ordenes_de_otra_empresa(self):
        from apps.empresas.models import Empresa

        otra = Empresa.objects.create(nombre_comercial='Otra', ruc='777777777777')
        ajena = self.OrdenTrabajo.objects.create(
            empresa=otra, cliente=self.cliente, vehiculo=self.vehiculo, numero_orden='OT-AJENA',
        )
        self.assertEqual(self.client.get(self._url(ajena)).status_code, 404)

class IconosSvgPdfTests(SimpleTestCase):
    """Los testigos del PDF usan los mismos paths SVG que la web."""

    def test_todos_los_testigos_se_interpretan(self):
        from apps.core.utils.svg_path import segmentos_svg

        from .testigos import TESTIGOS

        campos = {f.name for f in InspeccionVehiculo._meta.fields if f.name.startswith('testigo_')}
        self.assertEqual({t[0] for t in TESTIGOS}, campos)
        for campo, _, color, path in TESTIGOS:
            with self.subTest(campo=campo):
                segmentos = segmentos_svg(path)
                self.assertIn(color, ('red', 'yellow', 'green'))
                self.assertEqual(segmentos[0][0], 'M')
                for segmento in segmentos:
                    for valor in segmento[1:]:
                        self.assertTrue(-1 <= valor <= 25, (campo, segmento))

    def test_arcos_con_banderas_pegadas_y_comandos_relativos(self):
        from apps.core.utils.svg_path import segmentos_svg

        segmentos = segmentos_svg('M0 0a1 1 0 011 1h2v-1zM5,5Q6,6 7,5T9,5')
        self.assertEqual(segmentos[1][0], 'C')
        self.assertAlmostEqual(segmentos[1][-2], 1)
        self.assertAlmostEqual(segmentos[1][-1], 1)
        self.assertEqual(segmentos[2], ('L', 3, 1))
        self.assertEqual(segmentos[3], ('L', 3, 0))
        self.assertEqual(segmentos[4], ('Z',))
        self.assertEqual(segmentos[-1][-2:], (9.0, 5.0))

    def test_codigos_dtc_normalizados(self):
        from .pdf_inspeccion import codigos_dtc

        self.assertEqual(codigos_dtc(' p0301, P0420;p0301\nP0171 '), ['P0301', 'P0420', 'P0171'])
        self.assertEqual(codigos_dtc(None), [])


class NuevaHojaPdfTests(SimpleTestCase):
    """El plan de trabajo abre hoja nueva sin dejar hojas en blanco."""

    def _paginas(self, alto_relleno):
        from io import BytesIO

        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        from reportlab.lib.styles import getSampleStyleSheet

        from .pdf_inspeccion import _NuevaHoja

        doc = SimpleDocTemplate(BytesIO(), pagesize=A4)
        estilo = getSampleStyleSheet()['Normal']
        doc.build([Spacer(1, alto_relleno), _NuevaHoja(), Paragraph('Plan', estilo)])
        return doc.page

    def test_salta_si_la_hoja_tiene_contenido(self):
        self.assertEqual(self._paginas(100), 2)

    def test_no_deja_hoja_en_blanco_si_la_anterior_quedo_llena(self):
        from reportlab.lib.pagesizes import A4

        alto_marco = A4[1] - 2 * 72 - 12  # márgenes por defecto y padding del marco
        self.assertEqual(self._paginas(alto_marco), 2)
