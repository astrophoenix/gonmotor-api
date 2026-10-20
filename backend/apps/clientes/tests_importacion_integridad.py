"""Integridad de la importación de clientes y vehículos (fase 0).

Cada test reproduce un problema del análisis de importación masiva: la
importación no puede ser una puerta lateral de permisos, ni borrar datos con
celdas vacías, ni dejar dos dueños actuales, ni reactivar registros en
silencio, ni validar menos que la pantalla, ni escribir fórmulas en un Excel.
"""

from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient, APITestCase

from apps.authentication.models import UsuarioEmpresa
from apps.clientes.excel_import import CLIENT_COLUMNS
from apps.clientes.models import Cliente
from apps.core.testing import fijar_empresa
from apps.empresas.models import Empresa, Taller
from apps.vehiculos.models import Vehiculo, VehiculoEmpresa, VehiculoPropietario, vincular_empresa

URL = '/api/clientes/importar/'


def _xlsx(filas):
    columnas = [header for _, header in CLIENT_COLUMNS]
    claves = {header: key for key, header in CLIENT_COLUMNS}
    wb = Workbook()
    ws = wb.active
    ws.append(columnas)
    for fila in filas:
        ws.append([fila.get(claves[c], '') for c in columnas])
    buffer = BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile('clientes.xlsx', buffer.getvalue())


class _Base(APITestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre_comercial='Taller Integridad', ruc='1750000000012')
        self.taller = Taller.objects.create(empresa=self.empresa, nombre='Central', codigo_sucursal='001', direccion='x')
        self.admin = User.objects.create_superuser(username='imp-admin', password='x')
        self.client = self._cliente(self.admin)

    def _cliente(self, user):
        cliente = APIClient()
        cliente.force_authenticate(user=user)
        fijar_empresa(cliente, str(self.empresa.uuid))
        return cliente

    def _importar(self, filas, cliente=None, validar=False):
        url = URL + ('?validar=1' if validar else '')
        return (cliente or self.client).post(url, {'archivo': _xlsx(filas)}, format='multipart')

    def _vehiculo(self, placa, **datos):
        vehiculo = Vehiculo.objects.create(placa=placa, marca='Kia', modelo='Rio', **datos)
        vincular_empresa(vehiculo, self.empresa.pk)
        return vehiculo


class PermisosImportacionTests(_Base):

    def test_quien_no_modifica_vehiculos_no_los_crea_ni_los_edita(self):
        cajero = User.objects.create_user(username='imp-cajero', password='x')
        relacion = UsuarioEmpresa.objects.create(user=cajero, empresa=self.empresa, rol='CAJERO')
        relacion.talleres.set([self.taller])
        existente = self._vehiculo('PEX0001', color='Rojo')

        respuesta = self._importar([
            {'identificacion': '1711111111', 'nombre': 'Cliente Caja', 'placa': 'PNU0001', 'marca': 'Kia'},
            {'identificacion': '1711111111', 'nombre': 'Cliente Caja', 'placa': 'PEX0001', 'color': 'Azul'},
        ], cliente=self._cliente(cajero))

        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        # El cliente sí se importa (tiene permiso), los vehículos no.
        self.assertTrue(Cliente.objects.filter(empresa=self.empresa, identificacion='1711111111').exists())
        self.assertFalse(Vehiculo.objects.filter(placa='PNU0001').exists())
        existente.refresh_from_db()
        self.assertEqual(existente.color, 'Rojo')
        motivos = [e['motivo'] for e in respuesta.json()['errores']]
        self.assertEqual(len(motivos), 2)
        self.assertTrue(all('vehículos' in m for m in motivos), motivos)


class CeldasVaciasTests(_Base):

    def test_no_borran_datos_del_cliente(self):
        Cliente.objects.create(
            empresa=self.empresa, identificacion='1722222222', nombre='Ana',
            email='ana@correo.ec', telefono='0991111111', direccion='Calle 1', tipo_identificacion='R',
        )
        self._importar([{'identificacion': '1722222222', 'nombre': 'Ana Torres'}])
        cliente = Cliente.objects.get(empresa=self.empresa, identificacion='1722222222')
        self.assertEqual(cliente.nombre, 'Ana Torres')
        self.assertEqual(cliente.email, 'ana@correo.ec')
        self.assertEqual(cliente.telefono, '0991111111')
        self.assertEqual(cliente.direccion, 'Calle 1')
        self.assertEqual(cliente.tipo_identificacion, 'R')

    def test_no_borran_datos_del_vehiculo_ni_el_kilometraje(self):
        vehiculo = self._vehiculo('PVA0001', vin='VIN-ORIGINAL', color='Gris', kilometraje_actual=50000, anio=2018)
        self._importar([{'identificacion': '1733333333', 'nombre': 'Luis', 'placa': 'PVA0001'}])
        vehiculo.refresh_from_db()
        self.assertEqual(vehiculo.vin, 'VIN-ORIGINAL')
        self.assertEqual(vehiculo.color, 'Gris')
        self.assertEqual(vehiculo.anio, 2018)
        self.assertEqual(vehiculo.kilometraje_actual, 50000)

    def test_el_kilometraje_no_retrocede(self):
        vehiculo = self._vehiculo('PKM0001', kilometraje_actual=50000)
        respuesta = self._importar([
            {'identificacion': '1744444444', 'nombre': 'Rosa', 'placa': 'PKM0001', 'kilometraje_actual': '1000'},
        ])
        vehiculo.refresh_from_db()
        self.assertEqual(vehiculo.kilometraje_actual, 50000)
        self.assertIn('kilometraje', respuesta.json()['errores'][0]['motivo'].lower())

    def test_el_kilometraje_mayor_si_se_actualiza(self):
        vehiculo = self._vehiculo('PKM0002', kilometraje_actual=50000)
        self._importar([{'identificacion': '1744444445', 'nombre': 'Rosa', 'placa': 'PKM0002', 'kilometraje_actual': '61000'}])
        vehiculo.refresh_from_db()
        self.assertEqual(vehiculo.kilometraje_actual, 61000)


class PropietarioTests(_Base):

    def test_cambio_de_dueno_cierra_al_anterior(self):
        vehiculo = self._vehiculo('PDU0001')
        anterior = Cliente.objects.create(empresa=self.empresa, identificacion='1755555555', nombre='Dueño anterior')
        VehiculoPropietario.objects.create(vehiculo=vehiculo, cliente=anterior, es_actual=True)

        self._importar([{'identificacion': '1766666666', 'nombre': 'Dueño nuevo', 'placa': 'PDU0001'}])

        actuales = VehiculoPropietario.objects.filter(vehiculo=vehiculo, es_actual=True)
        self.assertEqual(actuales.count(), 1)
        self.assertEqual(actuales.get().cliente.identificacion, '1766666666')
        self.assertIsNotNone(VehiculoPropietario.objects.get(vehiculo=vehiculo, cliente=anterior).fecha_fin)


class DesactivadosTests(_Base):

    def test_cliente_desactivado_no_se_reactiva(self):
        Cliente.objects.create(empresa=self.empresa, identificacion='1777777777', nombre='Inactivo', is_active=False)
        respuesta = self._importar([{'identificacion': '1777777777', 'nombre': 'Inactivo'}])
        self.assertFalse(Cliente.objects.get(empresa=self.empresa, identificacion='1777777777').is_active)
        self.assertIn('desactivado', respuesta.json()['errores'][0]['motivo'])

    def test_vehiculo_suspendido_no_se_reactiva(self):
        vehiculo = self._vehiculo('PSU0001')
        VehiculoEmpresa.objects.filter(vehiculo=vehiculo, empresa=self.empresa).update(is_active=False)
        respuesta = self._importar([{'identificacion': '1788888888', 'nombre': 'Pedro', 'placa': 'PSU0001'}])
        self.assertFalse(VehiculoEmpresa.objects.get(vehiculo=vehiculo, empresa=self.empresa).is_active)
        self.assertIn('desactivado', respuesta.json()['errores'][0]['motivo'])


class ValidacionesTests(_Base):

    def test_email_invalido_se_rechaza(self):
        respuesta = self._importar([{'identificacion': '1799999990', 'nombre': 'Sin correo', 'email': 'no-es-correo'}])
        self.assertFalse(Cliente.objects.filter(identificacion='1799999990').exists())
        self.assertIn('correo', respuesta.json()['errores'][0]['motivo'].lower())

    def test_identificacion_se_normaliza_como_en_la_pantalla(self):
        Cliente.objects.create(empresa=self.empresa, identificacion='AB12345', nombre='Pasaporte', tipo_identificacion='P')
        respuesta = self._importar([{'identificacion': 'ab12345', 'nombre': 'Pasaporte Nuevo'}])
        self.assertEqual(respuesta.json()['actualizados'], 1)
        self.assertEqual(Cliente.objects.filter(empresa=self.empresa, identificacion__iexact='ab12345').count(), 1)

    def test_texto_demasiado_largo_da_un_mensaje_claro(self):
        respuesta = self._importar([{'identificacion': '1799999991', 'nombre': 'N' * 250}])
        motivo = respuesta.json()['errores'][0]['motivo']
        self.assertIn('Nombre', motivo)
        self.assertIn('200', motivo)
        self.assertNotIn('value too long', motivo)


class ArchivosExcelTests(_Base):

    def test_la_hoja_de_datos_de_la_plantilla_no_trae_filas_de_ejemplo(self):
        contenido = self.client.get(URL + 'plantilla/').content
        libro = load_workbook(BytesIO(contenido))
        datos = libro.worksheets[0]
        self.assertEqual(datos.max_row, 1)
        self.assertGreater(len(libro.worksheets), 1, 'Los ejemplos van en una hoja de instrucciones')

    def test_el_reporte_de_errores_no_escribe_formulas(self):
        contenido = self.client.post(URL + 'errores/', {'errores': [{
            'fila': 2, 'motivo': 'x', 'datos': {'nombre': '=HYPERLINK("http://malo","clic")'},
        }]}, format='json').content
        celdas = [c for fila in load_workbook(BytesIO(contenido)).active.iter_rows() for c in fila]
        # El texto llega intacto, pero como texto: no se evalúa al abrir el archivo.
        self.assertTrue(any(c.value == '=HYPERLINK("http://malo","clic")' for c in celdas))
        self.assertFalse([c.coordinate for c in celdas if c.data_type == 'f'])

    def test_la_exportacion_no_escribe_formulas(self):
        Cliente.objects.create(empresa=self.empresa, identificacion='1700000099', nombre='=1+1')
        contenido = self.client.get('/api/clientes/export-excel/').content
        celdas = [c for fila in load_workbook(BytesIO(contenido)).active.iter_rows() for c in fila]
        self.assertTrue(any(c.value == '=1+1' for c in celdas))
        self.assertFalse([c.coordinate for c in celdas if c.data_type == 'f'])
