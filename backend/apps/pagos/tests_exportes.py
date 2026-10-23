from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.utils import timezone
from openpyxl import load_workbook

from apps.caja.models import SesionCaja
from apps.ordenes.models import DetalleRepuestoOrdenTrabajo

from .models import PagoOrden
from .tests import ENDPOINT, PagosBase

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
URL_ORDENES = '/api/ordenes/ordenes-trabajo/export-excel/'
URL_PAGOS = f'{ENDPOINT}export-excel/'
URL_CXC = f'{ENDPOINT}cuentas-por-cobrar/export-excel/'


def filas_excel(respuesta, primera_columna):
    """Filas de datos de la tabla (las que siguen al encabezado `primera_columna`)."""
    hoja = load_workbook(BytesIO(respuesta.content)).active
    filas = list(hoja.iter_rows(values_only=True))
    inicio = next(i for i, fila in enumerate(filas) if fila[0] == primera_columna)
    return [fila for fila in filas[inicio + 1:] if any(valor not in (None, '') for valor in fila)]


class ExportesBase(PagosBase):

    def _excel(self, url, params=None, client=None, primera_columna=None):
        respuesta = (client or self._cliente(self.admin)).get(url, params or {})
        self.assertEqual(respuesta.status_code, 200, getattr(respuesta, 'content', b'')[:300])
        self.assertEqual(respuesta['Content-Type'], XLSX)
        return filas_excel(respuesta, primera_columna)


class ExportarOrdenesTests(ExportesBase):

    def setUp(self):
        super().setUp()
        self._pagar('15')
        self.orden_a2 = self._orden(self.empresa, self.taller_a2, self.cliente, 'A2-OT')

    def test_columnas_del_contador_con_cobro(self):
        filas = self._excel(URL_ORDENES, primera_columna='Nº Orden')
        por_numero = {fila[0]: fila for fila in filas}
        self.assertEqual(set(por_numero), {'A1-OT', 'A2-OT'})
        fila = por_numero['A1-OT']
        self.assertEqual(fila[1], 'A1')
        self.assertEqual(fila[6], '1700000011')
        # Base 0, gravada, descuento, IVA, total, pagado, saldo, estado de pago.
        self.assertEqual(
            [Decimal(str(valor)) for valor in fila[9:16]],
            [Decimal('0'), Decimal('100'), Decimal('0'), Decimal('15'), Decimal('115'), Decimal('15'), Decimal('100')],
        )
        self.assertEqual(fila[16], 'Pago parcial')

    def test_respeta_filtros_y_taller(self):
        filas = self._excel(URL_ORDENES, {'estado_pago': 'PARCIAL'}, primera_columna='Nº Orden')
        self.assertEqual([fila[0] for fila in filas], ['A1-OT'])
        # El cajero de A1 solo exporta su taller.
        filas = self._excel(URL_ORDENES, client=self.client, primera_columna='Nº Orden')
        self.assertEqual([fila[0] for fila in filas], ['A1-OT'])

    def test_otra_empresa_no_exporta_ordenes_ajenas(self):
        filas = self._excel(
            URL_ORDENES, client=self._cliente(self.cajero_ajeno, self.otra_empresa), primera_columna='Nº Orden',
        )
        self.assertEqual(filas, [])


class ExportarPagosTests(ExportesBase):

    def setUp(self):
        super().setUp()
        # La devolución en efectivo necesita efectivo en la caja (RF-04).
        SesionCaja.objects.filter(caja__taller=self.taller_a1, estado='ABIERTA').update(monto_inicial=Decimal('50'))
        self.pago = self._pagar('40', forma_pago='TRANSFERENCIA', referencia='=HYPERLINK("x")')
        self._pagar('10', tipo='DEVOLUCION')
        anulado = self._pagar('5')
        self.client.post(f'{ENDPOINT}{anulado.data["id"]}/anular/', {'motivo': 'Error de digitación'}, format='json')

    def test_devoluciones_en_negativo_y_anulados(self):
        filas = self._excel(URL_PAGOS, primera_columna='Fecha')
        montos = sorted(Decimal(str(fila[8])) for fila in filas)
        self.assertEqual(montos, [Decimal('-10'), Decimal('5'), Decimal('40')])
        anulado = next(fila for fila in filas if fila[11] == 'Sí')
        self.assertEqual(anulado[12], 'Error de digitación')
        self.assertEqual({fila[1] for fila in filas}, {'A1-OT'})

    def test_referencia_con_igual_no_es_formula(self):
        hoja = load_workbook(BytesIO(self._cliente(self.admin).get(URL_PAGOS).content)).active
        celdas = [celda for fila in hoja.iter_rows() for celda in fila if celda.value == '=HYPERLINK("x")']
        self.assertEqual(len(celdas), 1)
        self.assertEqual(celdas[0].data_type, 's')

    def test_filtros_del_listado(self):
        self.assertEqual(len(self._excel(URL_PAGOS, {'tipo': 'DEVOLUCION'}, primera_columna='Fecha')), 1)
        self.assertEqual(len(self._excel(URL_PAGOS, {'forma_pago': 'TRANSFERENCIA'}, primera_columna='Fecha')), 1)
        self.assertEqual(len(self._excel(URL_PAGOS, {'anulado': 'false'}, primera_columna='Fecha')), 2)
        # Rango de fechas: un pago de ayer queda fuera de «hoy».
        PagoOrden.objects.filter(uuid=self.pago.data['id']).update(fecha=timezone.now() - timedelta(days=1))
        hoy = timezone.localdate().isoformat()
        filas = self._excel(URL_PAGOS, {'fecha_desde': hoy, 'fecha_hasta': hoy}, primera_columna='Fecha')
        self.assertEqual(len(filas), 2)
        # El listado JSON usa los mismos filtros.
        self.assertEqual(self._cliente(self.admin).get(ENDPOINT, {'fecha_desde': hoy}).data['count'], 2)

    def test_permisos_y_alcance(self):
        self.assertEqual(self._cliente(self.asesor).get(URL_PAGOS).status_code, 403)
        ajeno = self._cliente(self.cajero_ajeno, self.otra_empresa)
        self.assertEqual(self._excel(URL_PAGOS, client=ajeno, primera_columna='Fecha'), [])


class ExportarCuentasPorCobrarTests(ExportesBase):

    def setUp(self):
        super().setUp()
        self._pagar('15')
        self.orden.estado = 'ENTREGADO'
        self.orden.fecha_entrega = timezone.now() - timedelta(days=3)
        self.orden.save()
        self.orden_a2 = self._orden(
            self.empresa, self.taller_a2, self.cliente, 'A2-OT',
            estado='ENTREGADO', fecha_entrega=timezone.now(),
        )
        DetalleRepuestoOrdenTrabajo.objects.create(orden_trabajo=self.orden_a2, descripcion='X', precio_unitario=Decimal('10'))
        self.orden_a2.calcular_totales()

    def test_mismas_filas_que_el_listado(self):
        filas = self._excel(URL_CXC, primera_columna='Orden')
        por_numero = {fila[0]: fila for fila in filas}
        self.assertEqual(set(por_numero), {'A1-OT', 'A2-OT'})
        fila = por_numero['A1-OT']
        self.assertEqual(fila[7], 3)
        self.assertEqual(Decimal(str(fila[10])), Decimal('100'))
        filas = self._excel(URL_CXC, {'search': 'A2-OT'}, primera_columna='Orden')
        self.assertEqual([fila[0] for fila in filas], ['A2-OT'])

    def test_cajero_solo_su_taller_y_asesor_sin_acceso(self):
        filas = self._excel(URL_CXC, client=self.client, primera_columna='Orden')
        self.assertEqual([fila[0] for fila in filas], ['A1-OT'])
        self.assertEqual(self._cliente(self.asesor).get(URL_CXC).status_code, 403)
