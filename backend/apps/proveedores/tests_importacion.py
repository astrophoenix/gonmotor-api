"""Importación de proveedores desde Excel."""

from io import BytesIO
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient, APITestCase

from apps.authentication.models import UsuarioEmpresa
from apps.core.testing import fijar_empresa
from apps.empresas.models import Empresa, Taller
from apps.proveedores.excel_import import ImportadorProveedores
from apps.proveedores.models import Proveedor

User = get_user_model()
URL = '/api/proveedores/importar/'


def _xlsx(filas):
    wb = Workbook()
    ws = wb.active
    ws.append([c.encabezado for c in ImportadorProveedores.columnas])
    for fila in filas:
        ws.append([fila.get(c.clave, '') for c in ImportadorProveedores.columnas])
    buffer = BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile('proveedores.xlsx', buffer.getvalue())


class ImportacionProveedoresTests(APITestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre_comercial='Taller Proveedores', ruc=str(uuid4().int % 10**13).zfill(13))
        self.taller = Taller.objects.create(empresa=self.empresa, nombre='Central', codigo_sucursal='001', direccion='x')
        self.client = self._como('ADMIN_EMPRESA')

    def _como(self, rol):
        user = User.objects.create_user(username=f'prov-{rol}-{uuid4().hex[:6]}', password='x')
        relacion = UsuarioEmpresa.objects.create(user=user, empresa=self.empresa, rol=rol)
        relacion.talleres.set([self.taller])
        cliente = APIClient()
        cliente.force_authenticate(user=user)
        fijar_empresa(cliente, str(self.empresa.uuid))
        return cliente

    def _importar(self, filas, validar=False, cliente=None):
        return (cliente or self.client).post(
            URL + ('?validar=1' if validar else ''), {'archivo': _xlsx(filas)}, format='multipart',
        )

    def test_crea_con_ruc_por_defecto(self):
        respuesta = self._importar([{'identificacion': '1790012345001', 'nombre': 'Repuestos del Valle'}])
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        proveedor = Proveedor.objects.get(empresa=self.empresa, identificacion='1790012345001')
        self.assertEqual(proveedor.tipo_identificacion, 'R')

    def test_existente_solo_cambian_las_celdas_con_valor(self):
        Proveedor.objects.create(
            empresa=self.empresa, identificacion='AB123', nombre='Proveedor', email='a@prov.ec', contacto='Ana',
        )
        respuesta = self._importar([{'identificacion': 'ab123', 'nombre': 'Proveedor Nuevo', 'telefono': '0990000000'}])
        self.assertEqual(respuesta.json()['actualizados'], 1)
        proveedor = Proveedor.objects.get(empresa=self.empresa, identificacion='AB123')
        self.assertEqual(proveedor.nombre, 'Proveedor Nuevo')
        self.assertEqual(proveedor.telefono, '0990000000')
        self.assertEqual(proveedor.email, 'a@prov.ec')
        self.assertEqual(proveedor.contacto, 'Ana')

    def test_validaciones(self):
        respuesta = self._importar([
            {'identificacion': '1790000000001', 'nombre': 'X', 'email': 'no-es-correo'},
            {'identificacion': '1790000000002', 'nombre': 'Y', 'tipo_identificacion': 'Z'},
            {'identificacion': '1790000000003', 'nombre': ''},
        ]).json()
        motivos = {e['fila']: e['motivo'] for e in respuesta['errores']}
        self.assertIn('correo', motivos[2].lower())
        self.assertIn('Tipo Identificación', motivos[3])
        self.assertIn('nombre', motivos[4])
        self.assertFalse(Proveedor.objects.filter(empresa=self.empresa).exists())

    def test_identificacion_repetida_y_desactivado(self):
        Proveedor.objects.create(empresa=self.empresa, identificacion='1790000000009', nombre='Inactivo', is_active=False)
        respuesta = self._importar([
            {'identificacion': '1790000000008', 'nombre': 'Uno'},
            {'identificacion': '1790000000008', 'nombre': 'Dos'},
            {'identificacion': '1790000000009', 'nombre': 'Inactivo'},
        ]).json()
        motivos = {e['fila']: e['motivo'] for e in respuesta['errores']}
        self.assertIn('fila 2', motivos[3])
        self.assertIn('desactivado', motivos[4])
        self.assertFalse(Proveedor.objects.get(empresa=self.empresa, identificacion='1790000000009').is_active)

    def test_vista_previa_no_guarda(self):
        previa = self._importar([{'identificacion': '1790000000010', 'nombre': 'Previa'}], validar=True).json()
        self.assertEqual(previa['filas'][0]['estado'], 'nuevo')
        self.assertEqual(previa['filas'][0]['identificacion'], '1790000000010')
        self.assertFalse(Proveedor.objects.filter(identificacion='1790000000010').exists())

    def test_solo_quien_modifica_proveedores_importa(self):
        # ASESOR: proveedores solo para ver.
        self.assertEqual(self._importar([{'identificacion': '1', 'nombre': 'X'}], cliente=self._como('ASESOR')).status_code, 403)

    def test_plantilla(self):
        libro = load_workbook(BytesIO(self.client.get(URL + 'plantilla/').content))
        self.assertEqual(libro.worksheets[0].max_row, 1)
        self.assertIn('Instrucciones', libro.sheetnames)
