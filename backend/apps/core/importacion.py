"""Núcleo de la importación masiva desde Excel (.xlsx).

Cada módulo define un `Importador` con sus columnas y cómo guardar sus filas;
este núcleo aporta lo común:

- Lectura del archivo: encabezados por nombre (sin importar tildes ni
  mayúsculas, con alias), columnas faltantes y límites de tamaño y de filas.
- Vista previa: con `validar=True` corre la misma lógica en una transacción que
  se revierte, así lo que se muestra es exactamente lo que se importará.
- Contrato de respuesta que consume `ImportWizardModal` del frontend: `total`,
  `exitosos`, `nuevos`, `actualizados`, `errores` y, al validar, `filas`.
- Motivos de error para el usuario (`FilaInvalida`); un error inesperado va al
  log y la fila recibe un mensaje genérico.
- Plantilla (hoja de datos solo con encabezados + listas desplegables, y hoja
  de instrucciones con ejemplos) y reporte de errores, sin fórmulas.
- Vistas genéricas para exponer las tres rutas de un módulo.

Las reglas de cada módulo (qué es un duplicado, qué permisos extra exige,
cómo se agrupa) viven en su importador. Referencia: `apps/clientes/excel_import.py`.
"""

import logging
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.conf import settings
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.renderers import JSONRenderer
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.authentication.utils import get_empresa_id_desde_request
from apps.core.utils.excel_export import neutralizar_formula

logger = logging.getLogger(__name__)

TIPO_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MENSAJE_ERROR_INESPERADO = 'No se pudo guardar la fila. Revisa sus datos e inténtalo de nuevo.'
FILAS_CON_LISTA = 5000

_BORDE = Border(
    left=Side(style='thin', color='e5e7eb'),
    right=Side(style='thin', color='e5e7eb'),
    top=Side(style='thin', color='e5e7eb'),
    bottom=Side(style='thin', color='e5e7eb'),
)
_ALINEACION_CABECERA = Alignment(horizontal='center', vertical='center', wrap_text=True)
_ALINEACION_CELDA = Alignment(vertical='center', wrap_text=True)


# ---------------------------------------------------------------------------
# Errores
# ---------------------------------------------------------------------------
class FilaInvalida(ValueError):
    """Motivo de rechazo de una fila, pensado para mostrarse tal cual al usuario."""


class ArchivoInvalido(ValueError):
    """El archivo completo no se puede procesar (formato, tamaño, filas)."""


class ColumnasFaltantesError(ArchivoInvalido):
    """El archivo no trae todas las columnas de la plantilla."""

    def __init__(self, columnas):
        self.columnas = columnas
        super().__init__(f'Faltan columnas en el archivo: {", ".join(columnas)}')


# ---------------------------------------------------------------------------
# Normalización de celdas
# ---------------------------------------------------------------------------
def limpiar(valor):
    """Normaliza un valor leído de Excel a texto."""
    if valor is None:
        return ''
    if isinstance(valor, bool):
        return 'true' if valor else 'false'
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def normalizar_encabezado(valor):
    texto = limpiar(valor).lower()
    texto = unicodedata.normalize('NFKD', texto)
    texto = ''.join(c for c in texto if not unicodedata.combining(c))
    texto = texto.replace('º', '')
    return ' '.join(texto.split())


def a_entero(valor):
    valor = limpiar(valor)
    if not valor:
        return None
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        raise ValueError(f'"{valor}" no es un número entero válido.')


def a_decimal(valor, etiqueta, decimales=2, maximo=Decimal('99999999.99')):
    """Decimal no negativo; acepta coma o punto decimal (12,50 o 12.50)."""
    valor = limpiar(valor)
    if not valor:
        return None
    texto = valor.replace(' ', '')
    if ',' in texto and '.' in texto:
        texto = texto.replace('.', '').replace(',', '.')  # 1.234,50 → 1234.50
    else:
        texto = texto.replace(',', '.')
    try:
        numero = Decimal(texto)
    except InvalidOperation:
        raise ValueError(f'{etiqueta}: "{valor}" no es un número válido.')
    if not numero.is_finite() or numero < 0:
        raise ValueError(f'{etiqueta}: debe ser un número mayor o igual a 0.')
    if numero > maximo:
        raise ValueError(f'{etiqueta}: el máximo es {maximo}.')
    return numero.quantize(Decimal(1).scaleb(-decimales))


def a_opcion(valor, choices, etiqueta):
    """Código de `choices` a partir del código o de la etiqueta (sin distinguir mayúsculas)."""
    valor = limpiar(valor).upper()
    if not valor:
        return None
    for codigo, texto in choices:
        if valor in (str(codigo).upper(), str(texto).upper()):
            return codigo
    raise ValueError(f'"{valor}" no es una opción válida para {etiqueta}.')


def validar_longitudes(modelo, datos, nombres):
    """Mensajes para los textos más largos que su columna en el modelo."""
    errores = []
    for campo, valor in datos.items():
        if not isinstance(valor, str):
            continue
        try:
            maximo = modelo._meta.get_field(campo).max_length
        except Exception:
            continue
        if maximo and len(valor) > maximo:
            errores.append(f'{nombres.get(campo, campo)}: máximo {maximo} caracteres (tiene {len(valor)}).')
    return errores


# ---------------------------------------------------------------------------
# Importador base
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Columna:
    clave: str
    encabezado: str
    alias: tuple = ()
    # Lista desplegable en la plantilla: choices de Django [(código, etiqueta)].
    opciones: tuple = field(default=())


class Importador:
    """Base de un importador. Una subclase define:

    - `columnas`: lista de `Columna` (el orden es el de la plantilla).
    - `hoja`, `entidad` (para mensajes) e `instrucciones`/`ejemplos` de la plantilla.
    - `procesar(filas)`: guarda las filas [(número, registro)] usando
      `self.registrar_error(...)`, `self.marcar(...)` y los contadores de
      `self.resultado`; cada escritura dentro de `transaction.atomic()`.
    - `resumen_fila(numero, registro)`: datos de la fila para la vista previa.
    """

    columnas = []
    hoja = 'Datos'
    entidad = 'registros'
    instrucciones = []
    ejemplos = []

    def __init__(self, empresa_id, request=None):
        self.empresa_id = empresa_id
        self.request = request

    # --- Contrato con las subclases -----------------------------------------
    def procesar(self, filas):
        raise NotImplementedError

    def resumen_fila(self, numero, registro):
        raise NotImplementedError

    def contadores_extra(self):
        """Contadores propios del módulo (p. ej. `vehiculos`), además de los comunes."""
        return {}

    # --- Ayudantes para `procesar` ------------------------------------------
    @property
    def nombres(self):
        return {c.clave: c.encabezado for c in self.columnas}

    def registrar_error(self, numero, identificador, registro, motivo):
        self.resultado['errores'].append({
            'fila': numero, 'identificacion': identificador, 'datos': registro, 'motivo': motivo,
        })
        self._estados[numero] = ('error', motivo)

    def marcar(self, numero, estado, motivo=''):
        self._estados[numero] = (estado, motivo)

    def motivo_de(self, error):
        # Solo los motivos pensados para el usuario; un error inesperado no
        # muestra textos internos (de la base de datos, por ejemplo).
        if isinstance(error, FilaInvalida):
            return str(error)
        logger.exception('Error inesperado al importar una fila de %s', self.entidad)
        return MENSAJE_ERROR_INESPERADO

    # --- Lectura --------------------------------------------------------------
    def leer(self, archivo):
        """Filas del archivo como [(número de fila, {clave: texto})]."""
        maximo_mb = settings.IMPORTACION_MAX_MB
        if getattr(archivo, 'size', 0) > maximo_mb * 1024 * 1024:
            raise ArchivoInvalido(f'El archivo supera el máximo de {maximo_mb} MB.')
        try:
            libro = load_workbook(archivo, read_only=True, data_only=True)
        except Exception as error:
            raise ArchivoInvalido('No se pudo abrir el archivo: verifica que sea un Excel .xlsx.') from error

        try:
            filas_hoja = libro.active.iter_rows(values_only=True)
            encabezados = next(filas_hoja, None)
            if encabezados is None:
                return []
            indices = self._mapear_columnas(encabezados)

            maximo_filas = settings.IMPORTACION_MAX_FILAS
            filas = []
            for numero, fila in enumerate(filas_hoja, start=2):
                if all(limpiar(celda) == '' for celda in fila):
                    continue
                if len(filas) >= maximo_filas:
                    raise ArchivoInvalido(
                        f'El archivo tiene más de {maximo_filas} filas con datos. '
                        'Divídelo en varios archivos e impórtalos por partes.'
                    )
                filas.append((numero, {
                    clave: limpiar(fila[indice]) if indice < len(fila) else ''
                    for clave, indice in indices.items()
                }))
            return filas
        finally:
            libro.close()

    def _mapear_columnas(self, encabezados):
        buscados = {}
        for columna in self.columnas:
            for nombre in (columna.encabezado, *columna.alias):
                buscados.setdefault(normalizar_encabezado(nombre), columna.clave)
        indices = {}
        for indice, encabezado in enumerate(encabezados):
            clave = buscados.get(normalizar_encabezado(encabezado))
            if clave and clave not in indices:
                indices[clave] = indice
        faltantes = [c.encabezado for c in self.columnas if c.clave not in indices]
        if faltantes:
            raise ColumnasFaltantesError(faltantes)
        return indices

    # --- Importación -----------------------------------------------------------
    def importar(self, archivo, validar=False):
        """Importa (o, con `validar`, simula sin guardar) y devuelve el resultado."""
        filas = self.leer(archivo)
        if not validar:
            return self._ejecutar(filas, validar=False)
        with transaction.atomic():
            resultado = self._ejecutar(filas, validar=True)
            transaction.set_rollback(True)
        return resultado

    def _ejecutar(self, filas, validar):
        self.resultado = {
            'total': len(filas), 'exitosos': 0, 'nuevos': 0, 'actualizados': 0,
            **self.contadores_extra(), 'errores': [],
        }
        self._estados = {}
        self.procesar(filas)
        if validar:
            self.resultado['filas'] = [
                {
                    'fila': numero,
                    **self.resumen_fila(numero, registro),
                    'estado': self._estados.get(numero, ('error', ''))[0],
                    'motivo': self._estados.get(numero, ('error', ''))[1],
                }
                for numero, registro in filas
            ]
        return self.resultado

    # --- Archivos que se descargan ---------------------------------------------
    def plantilla(self):
        """Hoja de datos (encabezados + listas) y hoja de instrucciones con ejemplos."""
        libro = Workbook()
        hoja = libro.active
        hoja.title = self.hoja
        relleno = PatternFill(start_color='1f2937', end_color='1f2937', fill_type='solid')
        for indice, columna in enumerate(self.columnas, start=1):
            celda = hoja.cell(row=1, column=indice, value=columna.encabezado)
            celda.font = Font(bold=True, color='FFFFFF', size=11)
            celda.fill = relleno
            celda.alignment = _ALINEACION_CABECERA
            celda.border = _BORDE
            hoja.column_dimensions[get_column_letter(indice)].width = 20
            if columna.opciones:
                letra = get_column_letter(indice)
                validacion = DataValidation(
                    type='list',
                    formula1='"' + ','.join(str(codigo) for codigo, _ in columna.opciones) + '"',
                    allow_blank=True,
                    showErrorMessage=True,
                    errorTitle='Valor no válido',
                    error='Elige un valor de la lista (ver hoja Instrucciones).',
                )
                validacion.add(f'{letra}2:{letra}{FILAS_CON_LISTA}')
                hoja.add_data_validation(validacion)
        hoja.freeze_panes = 'A2'

        ayuda = libro.create_sheet('Instrucciones')
        fila = 1
        ayuda.cell(row=fila, column=1, value=f'Cómo llenar la plantilla de {self.entidad}').font = Font(bold=True, size=13)
        fila += 2
        lineas = [
            f'Llena la hoja «{self.hoja}» desde la fila 2. No cambies los encabezados.',
            *self.instrucciones,
            f'Máximo {settings.IMPORTACION_MAX_FILAS} filas y {settings.IMPORTACION_MAX_MB} MB por archivo.',
        ]
        for linea in lineas:
            ayuda.cell(row=fila, column=1, value=f'• {linea}')
            fila += 1

        con_opciones = [c for c in self.columnas if c.opciones]
        if con_opciones:
            fila += 1
            ayuda.cell(row=fila, column=1, value='Códigos de las columnas con lista').font = Font(bold=True, size=11)
            fila += 1
            for columna in con_opciones:
                detalle = ', '.join(f'{codigo} = {texto}' for codigo, texto in columna.opciones)
                ayuda.cell(row=fila, column=1, value=f'{columna.encabezado}: {detalle}')
                fila += 1

        if self.ejemplos:
            fila += 1
            ayuda.cell(row=fila, column=1, value='Ejemplo (no lo copies a la hoja de datos tal cual)').font = Font(bold=True, size=11)
            fila += 1
            for indice, columna in enumerate(self.columnas, start=1):
                ayuda.cell(row=fila, column=indice, value=columna.encabezado).font = Font(bold=True)
            amarillo = PatternFill(start_color='fef9c3', end_color='fef9c3', fill_type='solid')
            for ejemplo in self.ejemplos:
                fila += 1
                for indice, valor in enumerate(ejemplo, start=1):
                    ayuda.cell(row=fila, column=indice, value=valor or '').fill = amarillo
        ayuda.column_dimensions['A'].width = 24
        return _a_buffer(libro)

    def reporte_errores(self, errores):
        """Excel con cada fila rechazada, su motivo y los datos tal como llegaron."""
        libro = Workbook()
        hoja = libro.active
        hoja.title = 'Errores de Importación'
        fecha = timezone.localtime().strftime('%d/%m/%Y %H:%M')
        titulo = hoja.cell(row=1, column=1, value=f'Reporte de errores de importación - Generado: {fecha}')
        titulo.font = Font(bold=True, size=14)
        encabezados = ['Fila', 'Motivo del Fallo'] + [c.encabezado for c in self.columnas]
        hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(encabezados))

        rojo = PatternFill(start_color='b91c1c', end_color='b91c1c', fill_type='solid')
        for indice, texto in enumerate(encabezados, start=1):
            celda = hoja.cell(row=3, column=indice, value=texto)
            celda.font = Font(bold=True, color='FFFFFF', size=11)
            celda.fill = rojo
            celda.alignment = _ALINEACION_CABECERA
            celda.border = _BORDE
            hoja.column_dimensions[get_column_letter(indice)].width = 22

        fila = 4
        for error in errores:
            datos = error.get('datos') or {}
            valores = [error.get('fila') or '', error.get('motivo') or '']
            valores += [datos.get(c.clave) or '' for c in self.columnas]
            for indice, valor in enumerate(valores, start=1):
                # Los datos vienen del Excel del usuario: nunca se escriben como fórmula.
                celda = neutralizar_formula(hoja.cell(row=fila, column=indice, value=valor))
                celda.alignment = _ALINEACION_CELDA
                celda.border = _BORDE
            fila += 1
        return _a_buffer(libro)


class ImportadorPorClave(Importador):
    """Registro de la empresa identificado por una clave única (código, identificación…).

    Reglas comunes (las mismas de la pantalla):

    - La clave se guarda en mayúsculas y define el duplicado; una clave repetida
      dentro del archivo es error (la segunda fila pisaría a la primera).
    - Si la clave existe, solo cambian las celdas con valor; un registro
      desactivado no se reactiva (la fila sale con su motivo).

    Una subclase define `modelo`, `columnas`, `campo_clave` (+ `nombre_clave`
    para los mensajes), `obligatorias`, `tipos` ({campo: 'decimal' | 'entero' |
    'email' | choices | tipo propio vía `convertir`}) y `por_defecto` (al crear).
    """

    modelo = None
    campo_clave = 'codigo'
    nombre_clave = 'el código'
    obligatorias = ()
    tipos = {}
    por_defecto = {}

    def resumen_fila(self, numero, registro):
        return {
            self.campo_clave: limpiar(registro.get(self.campo_clave)).upper(),
            'nombre': limpiar(registro.get('nombre')),
        }

    def convertir(self, tipo, valor, columna):
        """Tipos propios de un módulo (p. ej. el IVA del inventario)."""
        raise ValueError(f'{columna.encabezado}: tipo de dato no soportado.')

    def validar(self, registro):
        """Datos normalizados de la fila: obligatorios + opcionales con valor."""
        errores = []
        datos = {}
        for columna in self.columnas:
            valor = limpiar(registro.get(columna.clave))
            if not valor:
                if columna.clave in self.obligatorias:
                    errores.append(f'Falta {columna.encabezado.lower()}.')
                continue
            tipo = self.tipos.get(columna.clave, 'texto')
            try:
                if tipo == 'texto':
                    datos[columna.clave] = valor
                elif tipo == 'decimal':
                    datos[columna.clave] = a_decimal(valor, columna.encabezado)
                elif tipo == 'entero':
                    entero = a_entero(valor)
                    if entero < 0:
                        raise ValueError(f'{columna.encabezado}: debe ser mayor o igual a 0.')
                    datos[columna.clave] = entero
                elif tipo == 'email':
                    try:
                        validate_email(valor)
                    except DjangoValidationError:
                        raise ValueError(f'El correo electrónico "{valor}" no es válido.')
                    datos[columna.clave] = valor
                elif isinstance(tipo, (list, tuple)):
                    datos[columna.clave] = a_opcion(valor, tipo, columna.encabezado)
                else:
                    datos[columna.clave] = self.convertir(tipo, valor, columna)
            except ValueError as error:
                errores.append(str(error))
        if self.campo_clave in datos:
            datos[self.campo_clave] = datos[self.campo_clave].upper()
        errores += validar_longitudes(self.modelo, datos, self.nombres)
        if errores:
            raise FilaInvalida(' '.join(errores))
        return datos

    def ajustar_existente(self, existente, datos):
        """Último ajuste antes de actualizar un registro existente; devuelve un aviso o ''."""
        return ''

    def guardar(self, datos):
        """Crea o actualiza por la clave. Devuelve (creado, aviso)."""
        clave = datos[self.campo_clave]
        existente = self.modelo.objects.filter(
            empresa_id=self.empresa_id, **{self.campo_clave: clave},
        ).first()
        if existente is None:
            self.modelo.objects.create(empresa_id=self.empresa_id, **{**self.por_defecto, **datos})
            return True, ''
        if not existente.is_active:
            raise FilaInvalida(
                f'Ya existe un registro desactivado con {self.nombre_clave} {clave}. '
                'Actívalo desde el listado para poder actualizarlo.'
            )
        aviso = self.ajustar_existente(existente, datos)
        for campo, valor in datos.items():
            setattr(existente, campo, valor)
        existente.save(update_fields=[*datos.keys(), 'updated_at'])
        return False, aviso

    def procesar(self, filas):
        vistos = {}
        for numero, registro in filas:
            clave = limpiar(registro.get(self.campo_clave)).upper()
            try:
                if clave and clave in vistos:
                    raise FilaInvalida(
                        f'{self.nombre_clave.capitalize()} {clave} ya aparece en la fila {vistos[clave]} del archivo.'
                    )
                datos = self.validar(registro)
                vistos[clave] = numero
                with transaction.atomic():
                    creado, aviso = self.guardar(datos)
            except Exception as error:
                self.registrar_error(numero, clave, registro, self.motivo_de(error))
                continue
            self.marcar(numero, 'nuevo' if creado else 'actualiza', aviso)
            self.resultado['exitosos'] += 1
            self.resultado['nuevos' if creado else 'actualizados'] += 1


def _a_buffer(libro):
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer


def _respuesta_xlsx(buffer, nombre):
    respuesta = HttpResponse(buffer.read(), content_type=TIPO_XLSX)
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}"'
    return respuesta


# ---------------------------------------------------------------------------
# Vistas genéricas: un módulo solo indica su importador, permisos y nombres.
# ---------------------------------------------------------------------------
class _ImportacionBase(APIView):
    importador_class = None
    renderer_classes = [JSONRenderer]

    def importador(self, request):
        return self.importador_class(get_empresa_id_desde_request(request), request)


class ImportarView(_ImportacionBase):
    """`POST …/importar/[?validar=1]` con el archivo en `archivo`."""

    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        empresa_id = get_empresa_id_desde_request(request)
        if not empresa_id:
            return Response({'detail': 'No se pudo determinar la empresa activa.'}, status=403)

        archivo = request.FILES.get('archivo')
        if not archivo:
            return Response({'detail': 'No se envió ningún archivo. Debes adjuntar un archivo .xlsx.'}, status=400)
        if not (archivo.name or '').lower().endswith('.xlsx'):
            return Response({'detail': 'El archivo debe tener extensión .xlsx.'}, status=400)

        # ?validar=1 → vista previa: misma lógica, nada se guarda.
        validar = str(request.query_params.get('validar', '')).lower() in ('1', 'true', 'si', 'sí')
        try:
            resultado = self.importador(request).importar(archivo, validar=validar)
        except ColumnasFaltantesError as error:
            return Response({'detail': str(error), 'columnas_faltantes': error.columnas}, status=400)
        except ArchivoInvalido as error:
            return Response({'detail': str(error)}, status=400)
        except Exception:
            # El detalle técnico va al log, no al usuario.
            logger.exception('No se pudo procesar un archivo de importación (%s)', self.importador_class.__name__)
            return Response(
                {'detail': 'No se pudo procesar el archivo. Verifica que sea la plantilla en formato .xlsx.'},
                status=400,
            )
        estado = 400 if resultado['errores'] and not resultado['exitosos'] and not validar else 200
        return Response(resultado, status=estado)


class ErroresImportacionView(_ImportacionBase):
    """`POST …/importar/errores/` con `errores` (los de la respuesta) → Excel de errores."""

    nombre_archivo = 'errores_importacion.xlsx'

    def post(self, request):
        errores = request.data.get('errores')
        if not isinstance(errores, list):
            return Response({'detail': 'El campo "errores" debe ser una lista.'}, status=400)
        return _respuesta_xlsx(self.importador(request).reporte_errores(errores), self.nombre_archivo)


class PlantillaImportacionView(_ImportacionBase):
    """`GET …/importar/plantilla/` → plantilla oficial del módulo."""

    nombre_archivo = 'plantilla_importacion.xlsx'

    def get(self, request):
        return _respuesta_xlsx(self.importador(request).plantilla(), self.nombre_archivo)
