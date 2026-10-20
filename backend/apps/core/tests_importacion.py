"""Núcleo de importación masiva (`apps/core/importacion.py`): límites y archivos inválidos.

Se prueba a través de la importación de clientes, que es su primer usuario.
"""

from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient, APITestCase

from apps.clientes.excel_import import CLIENT_COLUMNS
from apps.clientes.models import Cliente
from apps.core.testing import fijar_empresa
from apps.empresas.models import Empresa

URL = '/api/clientes/importar/'


def _xlsx(n_filas):
    wb = Workbook()
    ws = wb.active
    ws.append([header for _, header in CLIENT_COLUMNS])
    for indice in range(n_filas):
        ws.append(['C', f'17{indice:08d}', f'Cliente {indice}'] + [''] * (len(CLIENT_COLUMNS) - 3))
    buffer = BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile('clientes.xlsx', buffer.getvalue())


class NucleoImportacionTests(APITestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre_comercial='Taller Núcleo', ruc='1750000000020')
        self.client = APIClient()
        self.client.force_authenticate(user=User.objects.create_superuser(username='nucleo', password='x'))
        fijar_empresa(self.client, str(self.empresa.uuid))

    @override_settings(IMPORTACION_MAX_FILAS=3)
    def test_rechaza_archivos_con_demasiadas_filas_sin_guardar_nada(self):
        respuesta = self.client.post(URL, {'archivo': _xlsx(4)}, format='multipart')
        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('más de 3 filas', respuesta.json()['detail'])
        self.assertFalse(Cliente.objects.filter(empresa=self.empresa).exists())

    @override_settings(IMPORTACION_MAX_FILAS=3)
    def test_acepta_hasta_el_limite(self):
        respuesta = self.client.post(URL, {'archivo': _xlsx(3)}, format='multipart')
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        self.assertEqual(respuesta.json()['nuevos'], 3)

    @override_settings(IMPORTACION_MAX_MB=0)
    def test_rechaza_archivos_demasiado_grandes(self):
        respuesta = self.client.post(URL, {'archivo': _xlsx(1)}, format='multipart')
        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('MB', respuesta.json()['detail'])

    def test_un_archivo_que_no_es_excel_da_un_mensaje_claro(self):
        falso = SimpleUploadedFile('clientes.xlsx', b'esto no es un excel')
        respuesta = self.client.post(URL, {'archivo': falso}, format='multipart')
        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('.xlsx', respuesta.json()['detail'])
        self.assertNotIn('zip', respuesta.json()['detail'].lower())

    def test_la_plantilla_indica_los_limites(self):
        contenido = self.client.get(URL + 'plantilla/').content
        instrucciones = load_workbook(BytesIO(contenido))['Instrucciones']
        textos = ' '.join(str(c.value) for fila in instrucciones.iter_rows() for c in fila if c.value)
        self.assertIn('Máximo 2000 filas', textos)
